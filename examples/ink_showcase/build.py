"""水墨 × 科研 —— 典型页面示范：封面 / 目录 / 章节过渡 / 展示页 ×4 / 结束页。

Run:   python build.py            -> out/水墨科研_示范.pptx + out/水墨科研_示范_preview/
图片： imagegen 生成的图按 IMAGES 里的文件名放进 img/，重新运行即自动替换程序画的占位图。

用到的“高级手法”
  · 书法字矢量化（brush.py）：书法字体转原生形状，换电脑不丢字体，可渐变 / 图片填充 / 错落排版
  · 图片填充文字：章节页巨型“壹”内透出墨韵纹理
  · 平滑切换 Morph：目录页的“壹”与章节名用同名 !!对象，切到章节页时放大飞入
  · 渐隐融图：封面 / 结束页的山水图左侧或上方渐隐进宣纸底，无硬边
  · 拱形窗图 + 双线框、拦腰大图 + 渐变蒙版 + 压边卡片、图片卡片上下双蒙版
  · 第 5、6 页直接用 skill 推荐版式：band_cards(topline 公式卡)、chart_page(原生折线图 + 关键发现)，
    通过替换 d.page 让它们套上水墨页眉页脚；折线图数据来自真实 MOC 算例，可在 PowerPoint 中编辑
"""
import os
import sys
import types

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'scripts'))
sys.path.insert(0, HERE)

import numpy as np                                           # noqa: E402
from lxml import etree                                       # noqa: E402
from pptstudio import Deck, Box                              # noqa: E402
from pptstudio import images as IMG, oxml as X              # noqa: E402
from brush import brush, seal, stroke                        # noqa: E402
import ink_assets as IA                                      # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 and sys.argv[1].endswith('.pptx') else \
    os.path.join(HERE, 'out', '水墨科研_示范.pptx')
W, H = 13.333, 7.5
PAPER, INK, DAI, DAI2, ZHU = 'F7F4EE', '14212B', '1E4A6B', '5C8AA6', 'B23A2E'   # 宣纸 墨 黛青 浅黛 朱砂

# ------------------------------------------------------------------ 图片（img/ 下有同名文件就用它）
IMAGES = {}


def img(name, fallback):
    p, real = IA.pick(name, fallback)
    IMAGES[name] = real
    return p


COVER = img('cover', lambda: IA.mountains((2560, 1440), 4, empty='left'))
AGENDA = img('agenda', lambda: IA.toned((1100, 1700), 21))
INKTEX = img('ink_texture', IA.ink_texture)
BAND = img('band', lambda: IA.toned((2560, 900), 12))
INNOV = [img(f'innov_{i}', lambda i=i: IA.toned((1200, 1600), 30 + i)) for i in (1, 2, 3)]
ENDING = img('ending', lambda: IA.mountains((2560, 1100), 7, empty='center'))
PAPER_IMG = IA.paper(PAPER)
SPLASH, _ = IA.pick('ink_splash', lambda: None)            # 可选：白底水墨晕染，垫在封面书法字后

SEC = ['研究背景与意义', '水锤机理与数值模型', '泄漏智能诊断方法', '工程应用与展望']
SEC_EN = ['BACKGROUND', 'MECHANISM & MODEL', 'INTELLIGENT DIAGNOSIS', 'APPLICATION & OUTLOOK']
NUM = '壹贰叁肆'


# ------------------------------------------------------------------ 公共部件
def base(d):
    cv = d.new(bg=False)
    cv.image(PAPER_IMG, 0, 0, W, H, fit='stretch', name='bg:paper')
    return cv


def arch(shape, top=50000, bottom=0):
    """把图片 / 形状改成拱形（上圆下方，月洞门的意思）。"""
    sp = X.spPr(shape)
    pg = sp.find(X.qn('a:prstGeom'))
    if pg is None:
        pg = X.insert_ordered(sp, etree.Element(X.qn('a:prstGeom')), X.SPPR_ORDER)
    pg.set('prst', 'round2SameRect')
    X.remove_children(pg, ['avLst'])
    av = X.sub(pg, 'a:avLst')
    X.sub(av, 'a:gd', name='adj1', fmla=f'val {top}')
    X.sub(av, 'a:gd', name='adj2', fmla=f'val {bottom}')
    return shape


def chrome(d, title, en=None):
    """内容页页眉页脚：朱砂章节印 + 宋体标题 + 字距英文；右上章节名；左下页码（章节取 d.cur_section）。"""
    cv = base(d)
    sec, page = max(0, d.cur_section), len(d.prs.slides)
    x, y = 0.75, 0.48
    seal(cv, x + 0.27, y + 0.33, 0.54, NUM[sec], font='brush', color=ZHU, seed=sec + 11, name='deco:secseal')
    tx = x + 0.78
    cv.text(tx, y - 0.04, 7.4, 0.5, title, size=24, color=INK, role='serif', bold=True, anchor='middle',
            wrap=False, fit=True, min_size=18, name='txt:title')
    cv.text(tx, y + 0.46, 7.4, 0.24, en or '', size=8.5, color='muted', role='display', spacing=2.4, wrap=False,
            fit=True, min_size=7, name='txt:title_en')
    cv.text(W - 0.75 - 3.4, y + 0.02, 3.4, 0.3, f'{NUM[sec]}  ·  {SEC[sec]}', size=10, color='muted', align='right',
            wrap=False, fit=True, name='txt:runner')
    cv.line(W - 0.75 - 0.9, y + 0.4, W - 0.75, y + 0.4, 'line', 0.75)
    cv.text(0.75, H - 0.5, 0.5, 0.25, f'{page:02d}', size=10, color='muted', role='display', spacing=1, wrap=False,
            fit=False, name='txt:page')
    cv.line(1.2, H - 0.375, 1.9, H - 0.375, 'line', 0.75)
    cv.content = Box(0.75, 1.55, W - 1.5, H - 1.55 - 0.6)
    return cv


def ink_page(self, title, en=None, source=None, bg=True, section=None, nav=True, num=None):
    """替换 Deck.page：skill 版式（band_cards / chart_page …）照常调用 d.page()，拿到的是水墨页眉页脚。"""
    if section is not None:
        self.cur_section = section - 1 if isinstance(section, int) else self.sections.index(section)
    cv = chrome(self, title, en)
    if source:
        b = cv.content
        cv.content = Box(b.x, b.y, b.w, b.h - 0.3)
        cv.text(b.x, cv.content.b + 0.08, b.w, 0.24, source, size=9, color='muted', wrap=False, fit=True,
                min_size=7, name='txt:source')
    self.pages[-1] = ('content', title)
    return cv


def _named(cv, *names):
    return [sh for sh in cv.slide.shapes if sh.name in names]


def _drop(shape):
    shape._element.getparent().remove(shape._element)


def _inches(shape):
    return shape.left / 914400, shape.top / 914400, shape.width / 914400, shape.height / 914400


def ink_band(cv):
    """band_cards 的蓝色渐变 + 点阵（科技感）→ 水墨色带：墨→黛青渐变，右侧叠一抹淡白远山；卡片顶线改朱砂。"""
    band = _named(cv, 'bg:band')[0]
    X.set_grad(band, [(0, INK, 1.0), (100, '24465E', 1.0)], 0)
    for sh in _named(cv, 'bg:banddots'):
        _drop(sh)
    x, y, w, h = _inches(band)
    mw = 7.2
    hills = IMG.cover_crop(ENDING, mw / h, focus=(0.72, 0.82))
    hills = IA.alpha_from_white(hills, 'band_hills.png', color='FFFFFF', strength=0.55)
    hills = IMG.fade(hills, side='left', start=0.35, end=0.95)
    pic = cv.image(hills, W - mw, y, mw, h, fit='stretch', name='bg:bandhills')
    band._element.addnext(pic._element)                    # 紧贴色带之上、文字之下
    for sh in _named(cv, 'deco:cardtop'):
        X.set_fill(sh, ZHU)
    serif = cv.t.f_serif                                   # 色带标题与页面标题同用宋体
    for sh in _named(cv, 'txt:bandhead'):
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                rpr = r._r.get_or_add_rPr()
                for tag in ('a:latin', 'a:ea'):
                    el = rpr.find(X.qn(tag))
                    if el is not None:
                        el.set('typeface', serif)


def ink_summary(cv, text):
    """chart_page 底部的渐变胶囊 → 宣纸色结论条 + 朱砂竖线（与第 4 页、目录页同一语汇）。"""
    x, y, w, h = _inches(_named(cv, 'deco:summary')[0])
    for sh in _named(cv, 'deco:summary', 'txt:summary', 'icon:lightbulb'):
        _drop(sh)
    cv.rect(x, y, w, h, fill='EFEBE3', name='deco:stmt')
    cv.rect(x, y, 0.06, h, fill=ZHU, name='deco:stmtbar')
    cv.text(x + 0.32, y, w - 0.5, h, text, size=16, color=INK, role='serif', bold=True, anchor='middle',
            wrap=False, fit=True, name='txt:stmt')


def card(cv, b, shadow=True, fill='FFFFFF'):
    return cv.rect(b.x, b.y, b.w, b.h, fill=fill, radius=0.04,
                   shadow={'color': INK, 'alpha': 0.11, 'blur': 26, 'dist': 6} if shadow else None, name='deco:card')


# ------------------------------------------------------------------ 1 封面
def cover(d):
    cv = base(d)
    pic = IMG.fade(IMG.cover_crop(COVER, W / H, focus=(0.6, 0.5)), side='left', start=0.45, end=0.8)
    cv.image(pic, 0, 0, W, H, fit='stretch', name='img:cover')
    cv.circle(10.35, 1.4, 0.32, fill=ZHU, alpha=0.86, name='deco:sun')                  # 一轮红日
    if SPLASH:
        cv.image(IA.alpha_from_white(SPLASH, 'splash_cover.png', color=DAI, strength=0.35), 0.2, 0.25, 6.6, 3.4,
                 fit='contain', name='img:splash')
    _, hb = brush(cv, '水脉', 0.82, 0.78, em=2.55, grad=[(0, INK, 1), (100, DAI, 1)], grad_angle=90,
                  name='brush:hero')
    # 题款（竖排行楷）+ 印章：中国画“题跋”构图
    _, vb = brush(cv, '千里引水一脉安澜', 7.0, 0.8, em=0.36, font='xingkai', vertical=True, spacing=0.12,
                  fill=INK, alpha=0.9, name='brush:inscription')
    seal(cv, vb.cx, vb.b + 0.38, 0.42, '安澜', color=ZHU, seed=5)
    ty = hb.b + 0.42
    cv.text(0.88, ty, 6.4, 1.25, ['长距离输水系统水锤防护', '与泄漏智能诊断研究'], size=30, color=INK, role='serif',
            bold=True, ls=1.12, sa=4, fit=True, name='txt:cover_title')
    cv.text(0.9, ty + 1.32, 4.6, 0.3, 'WATER HAMMER  ·  LEAK DIAGNOSIS', size=10,
            color='muted', role='display', spacing=2.2, wrap=False, fit=True, name='txt:cover_en')
    stroke(cv, 0.86, ty + 1.8, 1.35, 0.14, fill=ZHU, name='deco:cover_stroke')
    cv.text(0.9, ty + 2.1, 6.5, 0.7, ['汇报人：张三　　指导教师：李四 教授', 'XX大学 水利水电学院　·　2026 年 10 月'],
            size=12.5, color='text', ls=1.2, sa=6, fit=True, name='txt:cover_info')
    cv.transition('fade')
    cv.notes('封面：书法“水脉”为矢量形状（渐变墨色），右侧山水图左缘渐隐进宣纸底；竖排题款 + 印章取中国画题跋构图。')


# ------------------------------------------------------------------ 2 目录
DESC = ['工程规模、事故案例与研究缺口', '一维瞬变流方程与特征线解法', '压力波反射特征与深度学习定位', '示范工程验证与后续研究计划']


def agenda(d):
    cv = base(d)
    ax, ay, aw, ah = 0.95, 0.78, 3.3, 5.95
    o = 0.13
    sh = cv.shape('rect', ax - o, ay - o, aw + 2 * o, ah + 2 * o, fill=None, line=DAI, line_w=0.75, line_alpha=0.55,
                  name='deco:archline')
    arch(sh)
    pic = cv.image(AGENDA, ax, ay, aw, ah, fit='cover', focus=(0.5, 0.6), name='!!arch')
    arch(pic)
    mask = cv.rect(ax, ay, aw, ah, fill=None, grad=[(0, INK, 0.0), (50, INK, 0.08), (100, INK, 0.72)],
                   grad_angle=90, name='deco:archmask')
    arch(mask)
    _, mb = brush(cv, '目录', ax, ay + ah - 1.85, em=1.08, fill='FFFFFF', w=aw, align='center', name='brush:mulu')
    cv.text(ax, mb.b + 0.16, aw, 0.3, 'CONTENTS', size=10.5, color='FFFFFF', role='display', spacing=7,
            align='center', alpha=0.85, wrap=False, fit=False, name='txt:contents')
    seal(cv, ax + aw + o, ay + 1.75, 0.5, '安澜', color=ZHU, seed=8)                    # 压在框线上
    rx0, rx1 = 5.2, W - 0.8
    top, rh = 0.88, 1.43
    for i in range(4):
        y = top + i * rh
        brush(cv, NUM[i], rx0, y, em=0.74, fill=INK, h=rh, anchor='middle', name=f'!!num{i + 1}')
        tx = rx0 + 1.2
        cv.text(tx, y + rh / 2 - 0.5, 3.2, 0.5, SEC[i], size=22, color=INK, role='serif', bold=True,
                anchor='bottom', wrap=False, fit=True, name=f'!!sec{i + 1}')
        cv.text(tx, y + rh / 2 + 0.08, 3.2, 0.26, SEC_EN[i], size=10, color='muted', role='display', spacing=2.6,
                wrap=False, fit=True, name=f'txt:secen{i}')
        cv.text(rx1 - 2.9, y + rh / 2 - 0.3, 2.9, 0.6, DESC[i], size=11.5, color='muted', align='right',
                anchor='middle', fit=True, name=f'txt:desc{i}')
        if i < 3:
            cv.line(rx0, y + rh, rx1, y + rh, 'line', 0.75)
            cv.circle(rx0, y + rh, 0.035, fill=ZHU, name='deco:dot')
    cv.transition('fade')
    cv.notes('目录：拱形窗图（round2SameRect）+ 外双线框，印章压在框线上；序号用书法大写数字。'
             '“壹”与章节名对象名以 !! 开头，下一页 Morph 时会放大飞入。')


# ------------------------------------------------------------------ 3 章节过渡（Morph）
SEC_LONG_EN = ['Background & Significance', 'Mechanism & Numerical Model', 'Intelligent Leak Diagnosis',
               'Application & Outlook']


def section(d, i, bullets):
    cv = base(d)
    cv.circle(9.75, 3.78, 2.8, fill='EDE8DE', name='!!moon')
    cv.circle(9.75, 3.78, 2.98, fill=None, line=DAI, line_w=0.6, line_alpha=0.35, name='deco:moonline')
    brush(cv, NUM[i], 7.05, 0.85, em=5.7, image=INKTEX, w=5.4, h=5.85, align='center', anchor='middle',
          shadow={'color': INK, 'alpha': 0.2, 'blur': 20, 'dist': 7}, name=f'!!num{i + 1}')
    seal(cv, 12.55, 6.72, 0.44, '安澜', color=ZHU, seed=9)
    cv.text(0.95, 1.6, 4, 0.3, f'PART  0{i + 1}', size=12, color=ZHU, role='display', spacing=5, wrap=False,
            fit=False, name='txt:part')
    cv.text(0.95, 1.98, 6.3, 0.9, SEC[i], size=40, color=INK, role='serif', bold=True, anchor='middle', wrap=False,
            fit=True, name=f'!!sec{i + 1}')
    cv.text(0.95, 2.92, 6.0, 0.32, SEC_LONG_EN[i], size=13, color='muted', role='display', spacing=1.5,
            wrap=False, fit=True, name='txt:sec_en')
    stroke(cv, 0.92, 3.45, 1.6, 0.16, fill=ZHU)
    for k, t in enumerate(bullets):
        y = 4.0 + k * 0.56
        cv.text(0.95, y, 0.5, 0.36, f'0{k + 1}', size=11, color=ZHU, role='display', anchor='middle', wrap=False,
                fit=False, name='txt:bnum')
        cv.text(1.45, y, 5.4, 0.36, t, size=15, color='text', anchor='middle', wrap=False, fit=True,
                name='txt:bullet')
        cv.line(1.45, y + 0.46, 5.6, y + 0.46, 'line', 0.5)
    for j in range(4):
        cv.rect(0.95 + j * 0.62, 6.42, 0.5, 0.045, fill=ZHU if j == i else 'DAD4C8', name='deco:progress')
    cv.text(0.95 + 4 * 0.62 + 0.1, 6.3, 1.0, 0.28, f'{i + 1} / 4', size=10, color='muted', role='display',
            wrap=False, fit=False, name='txt:progress')
    cv.transition('morph')
    cv.notes('章节页：巨型“壹”是书法矢量形状 + 图片填充（墨韵纹理），与目录页同名对象 Morph 衔接。')


# ------------------------------------------------------------------ 4 展示页：拦腰大图 + 压边数据卡
CARDS = [
    ('1000', 'm/s', '水锤波速', '钢管中压力波速约 900–1200 m/s，数秒即可往返数公里管线'),
    ('122', 'm', '压力骤升', '流速骤减 1.2 m/s 时，按 Joukowsky 公式水头升高约 122 m'),
    ('60', 's', '相长 2L/a', '30 km 管线的相长约 60 s，关阀快于此即形成直接水锤'),
]


def page_background(d):
    d.cur_section = 0
    cv = chrome(d, '长距离输水的规模与风险', 'SCALE AND RISKS OF LONG-DISTANCE WATER TRANSFER')
    by, bh = 1.6, 2.6
    toned = IMG.duotone(BAND, dark='0E1A24', light='DCE4E6', contrast=1.05)   # 写实照片 → 墨色双色调
    cv.image(toned, 0, by, W, bh, fit='cover', focus=(0.5, 0.55), name='img:band')
    cv.rect(0, by, W, bh, fill=None, grad=[(0, INK, 0.86), (42, INK, 0.5), (100, INK, 0.0)], grad_angle=0,
            name='deco:bandmask')
    cv.text(0.95, by + 0.36, 6.6, 0.6, '一条输水线，就是一条“生命线”', size=24, color='FFFFFF', role='serif',
            bold=True, anchor='middle', wrap=False, fit=True, name='txt:band_title')
    cv.text(0.95, by + 1.04, 5.7, 0.75, '停泵、关阀等瞬变工况下，压力波在管道中高速往返；线路越长、流速越大，水锤风险越集中。',
            size=13, color='FFFFFF', alpha=0.9, ls=1.3, fit=True, name='txt:band_text')
    cy0 = by + bh - 0.48
    cols = Box(0.75, cy0, W - 1.5, 6.3 - cy0).cols(3, gap=0.32)
    for k, (c, (val, unit, label, desc)) in enumerate(zip(cols, CARDS)):
        card(cv, c)
        brush(cv, '一二三'[k], c.r - 1.25, c.y + 0.32, em=1.0, fill='EEE9E0', w=1.0, align='right',
              name='deco:cardnum')                                                       # 角落淡书法序号
        cv.rect(c.x + 0.36, c.y, 0.5, 0.05, fill=ZHU, name='deco:cardtop')
        vw = cv.width_of(val, 46, role='number') * 1.08
        cv.text(c.x + 0.34, c.y + 0.3, vw + 0.1, 0.85, val, size=46, color=INK, role='number', anchor='bottom',
                wrap=False, fit=False, name='txt:val')
        cv.text(c.x + 0.34 + vw + 0.1, c.y + 0.3, 1.2, 0.76, unit, size=16, color=DAI, role='number',
                anchor='bottom', wrap=False, fit=False, name='txt:unit')
        cv.text(c.x + 0.36, c.y + 1.25, c.w - 0.7, 0.4, label, size=15, color=INK, role='serif', bold=True,
                anchor='middle', wrap=False, fit=True, name='txt:label')
        cv.text(c.x + 0.36, c.y + 1.72, c.w - 0.72, c.h - 1.9, desc, size=12, color='text', ls=1.35, fit=True,
                name='txt:desc')
    cv.text(0.75, 6.5, 9, 0.24, '注：数值按钢管典型波速 a ≈ 1000 m/s 估算，仅作版式示例。', size=9, color='muted',
            wrap=False, fit=True, name='txt:source')
    cv.transition('fade')


# ------------------------------------------------------------------ 5 展示页：拦腰 + 公式卡（skill 版式 band_cards）
def page_moc(d):
    d.cur_section = 1
    cv = d.band_cards(
        '水锤控制方程与特征线解法', en='GOVERNING EQUATIONS & METHOD OF CHARACTERISTICS', card_style='topline',
        headline='以一维瞬变流方程描述水锤，沿特征线 [[dx/dt = ±a]] 化为常微分方程求解',
        text='取 Δt = Δx / a（Courant 数 = 1）时，特征线恰好穿过网格节点，无需插值、数值耗散最小',
        pills=['连续方程', '动量方程', '特征线法', 'Courant = 1'],
        cards=[{'title': '连续方程', 'num': 1,
                'formula': r'\dfrac{\partial H}{\partial t}+\dfrac{a^{2}}{gA}\,\dfrac{\partial Q}{\partial x}=0',
                'symbols': [('H', '测压管水头'), ('Q', '流量'), ('a', '水锤波速')]},
               {'title': '动量方程', 'num': 2,
                'formula': r'\dfrac{\partial Q}{\partial t}+gA\,\dfrac{\partial H}{\partial x}+\dfrac{f\,Q|Q|}{2DA}=0',
                'symbols': [('f', '摩阻系数'), ('D', '管径'), ('A', '断面面积')]},
               {'title': '特征线相容方程', 'num': 3,
                'formula': r'H_P=C_{P,M}\mp B\,Q_P',
                'symbols': [(r'C^{\pm}', '沿 dx/dt = ±a'), ('C_P,C_M', '上一时层已知'), ('B', 'a / (gA)')]}],
        notes="公式页：skill 推荐的 band_cards(card_style='topline')——色带写结论，卡片放公式 + 符号说明条。")
    ink_band(cv)


# ------------------------------------------------------------------ 6 展示页：原生折线图 + 关键发现（skill 版式 chart_page）
def moc_run(tc, L=3000.0, a=1000.0, D=0.8, f=0.018, Q0=0.45, HR=100.0, N=30, T=60.0, Em=1.5):
    """水库—管道—阀门系统，阀门按 τ = (1 − t/tc)^Em 关闭；返回 (t, 阀前水头, 稳态水头)。"""
    g = 9.81
    Ap = np.pi * D * D / 4
    dx = L / N
    dt = dx / a
    B, R = a / (g * Ap), f * dx / (2 * g * D * Ap * Ap)
    Hh = HR - np.arange(N + 1) * R * Q0 * Q0
    Q = np.full(N + 1, Q0)
    H0 = Hh[-1]
    ts, hv, t = [0.0], [H0], 0.0
    while t < T - 1e-9:
        t += dt
        CP = Hh[:-1] + B * Q[:-1] - R * Q[:-1] * np.abs(Q[:-1])
        CM = Hh[1:] - B * Q[1:] + R * Q[1:] * np.abs(Q[1:])
        HP, QP = np.empty_like(Hh), np.empty_like(Q)
        HP[1:-1] = (CP[:-1] + CM[1:]) / 2
        QP[1:-1] = (CP[:-1] - CM[1:]) / (2 * B)
        HP[0], QP[0] = HR, (HR - CM[0]) / B
        tau = (1 - t / tc) ** Em if t < tc else 0.0
        Cv = (Q0 * tau) ** 2 / (2 * H0)
        QP[-1] = -B * Cv + np.sqrt((B * Cv) ** 2 + 2 * Cv * CP[-1]) if Cv > 0 else 0.0
        HP[-1] = CP[-1] - B * QP[-1]
        Hh, Q = HP, QP
        ts.append(t)
        hv.append(Hh[-1])
    return np.array(ts), np.array(hv), H0


def tidy_timeseries(cv, label_every, dashed=()):
    """chart_page 的折线图默认带圆点标记、逐个类别标刻度；时程曲线改成：无标记、细线、隔 N 点标一次刻度。"""
    from pptx.enum.chart import XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.util import Pt
    ch = next(s for s in cv.slide.shapes if s.has_chart).chart
    for i, s in enumerate(ch.plots[0].series):
        s.marker.style = XL_MARKER_STYLE.NONE
        s.smooth = False
        s.format.line.width = Pt(1.0 if i in dashed else 1.75)
        if i in dashed:
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    c_ns = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    cat = ch._chartSpace.find('.//{%s}catAx' % c_ns)
    anchor = cat.find('{%s}noMultiLvlLbl' % c_ns)
    for tag in ('tickLblSkip', 'tickMarkSkip'):
        el = etree.Element('{%s}%s' % (c_ns, tag), val=str(label_every))
        if anchor is not None:
            anchor.addprevious(el)
        else:
            cat.append(el)


def page_result(d):
    d.cur_section = 1
    t1, h1, H0 = moc_run(3.0)
    _, h2, _ = moc_run(30.0)
    k = 2                                                     # 0.2 s 取一点
    cut = (1 - (h2.max() - H0) / (h1.max() - H0)) * 100
    cv = d.chart_page(
        '关阀方式对水锤压力的影响', en='VALVE CLOSURE VS. PRESSURE SURGE  ·  MOC CASE STUDY',
        chart={'kind': 'line_plain', 'categories': [f'{t:.0f}' for t in t1[::k]],
               'series': {'快速关阀  tc = 3 s': np.round(h1[::k], 1).tolist(),
                          '缓闭关阀  tc = 30 s': np.round(h2[::k], 1).tolist(),
                          '初始稳态水头': [round(H0, 1)] * len(t1[::k])},
               'colors': [ZHU, DAI, 'B5B0A6'], 'labels': False, 'legend': True, 'value_axis': True,
               'gridlines': True, 'number_format': '0', 'min_value': 0, 'max_value': 220},
        chart_title='图  阀前测压管水头时程（MOC 算例，示意）', unit='横轴 t / s　纵轴 H / m',
        points=[f'快关 tc = 3 s < 2L/a，形成[[直接水锤]]，峰值约 {h1.max():.0f} m',
                f'缓闭 tc = 30 s，峰值约 {h2.max():.0f} m，升压降低 [[{cut:.0f}%]]',
                '压力波以 4L/a = 12 s 为周期往复振荡'],
        summary='关阀历时大于 2L/a，可避免直接水锤',
        source='算例参数：L = 3 km，a = 1000 m/s，D = 0.8 m，Q₀ = 0.45 m³/s，f = 0.018；阀门按 τ = (1 − t/tc)^1.5 关闭。')
    tidy_timeseries(cv, label_every=int(round(10 / (0.1 * k))), dashed=(2,))
    ink_summary(cv, '关阀历时大于 [[2L/a]]，可避免直接水锤')
    print(f'MOC: H0={H0:.1f} m, max fast={h1.max():.1f} m, max slow={h2.max():.1f} m, cut={cut:.1f} %')


# ------------------------------------------------------------------ 7 展示页：图片卡片（创新点）
INNOV_TXT = [('MECHANISM', '多因素耦合水锤模型', '计入管壁黏弹性、气液两相与非恒定摩阻，刻画压力波的衰减与畸变'),
             ('ALGORITHM', '瞬变信号智能诊断', '以 MOC 仿真样本训练深度网络，从压力波反射中识别并定位泄漏'),
             ('ENGINEERING', '防护—诊断一体化', '阀门调度、空气阀布置与在线监测联动，形成闭环安全保障')]


def page_innov(d):
    d.cur_section = 2
    cv = chrome(d, '研究特色与创新', 'HIGHLIGHTS  &  INNOVATIONS')
    cols = Box(0.75, 1.62, W - 1.5, 5.0).cols(3, gap=0.3)
    for k, (c, img_p, (en, title, desc)) in enumerate(zip(cols, INNOV, INNOV_TXT)):
        cv.image(img_p, c.x, c.y, c.w, c.h, fit='cover', radius=0.06,
                 shadow={'color': INK, 'alpha': 0.16, 'blur': 22, 'dist': 6}, name=f'img:innov{k + 1}')
        cv.rect(c.x, c.y, c.w, c.h, fill=None, radius=0.06,
                grad=[(0, INK, 0.5), (26, INK, 0.0), (45, INK, 0.0), (100, INK, 0.9)], grad_angle=90,
                name='deco:innovmask')
        seal(cv, c.r - 0.48, c.y + 0.48, 0.5, '一二三'[k], font='brush', color=ZHU, seed=20 + k,
             name='deco:innovseal')                                                  # 朱砂小印作序号
        cv.text(c.x + 0.36, c.b - 2.1, c.w - 0.7, 0.26, en, size=10, color='FFFFFF', alpha=0.75, role='display',
                spacing=3, wrap=False, fit=True, name='txt:innov_en')
        cv.rect(c.x + 0.36, c.b - 1.74, 0.42, 0.045, fill=ZHU, name='deco:innovbar')
        cv.text(c.x + 0.36, c.b - 1.62, c.w - 0.7, 0.5, title, size=20, color='FFFFFF', role='serif', bold=True,
                anchor='middle', wrap=False, fit=True, name='txt:innov_title')
        cv.text(c.x + 0.36, c.b - 1.05, c.w - 0.7, 0.85, desc, size=12, color='FFFFFF', alpha=0.88, ls=1.35,
                fit=True, name='txt:innov_desc')
    cv.transition('fade')


# ------------------------------------------------------------------ 8 结束页
def ending(d):
    cv = base(d)
    hh = 4.4
    pic = IMG.fade(IMG.cover_crop(ENDING, W / hh, focus=(0.5, 0.6)), side='top', start=0.3, end=0.95)
    cv.image(pic, 0, H - hh, W, hh, fit='stretch', name='img:ending')
    _, b = brush(cv, '谢谢', 0, 1.15, em=2.45, grad=[(0, INK, 1), (100, DAI, 1)], grad_angle=90, w=W,
                 align='center', name='brush:thanks')
    seal(cv, b.r + 0.42, b.b - 0.38, 0.6, '安澜', color=ZHU, seed=5)
    cv.transition('fade')


# ------------------------------------------------------------------ build
def build():
    d = Deck('academic_navy', primary=DAI, secondary=DAI2, accent=ZHU, emph=ZHU, dark=INK, text='3A434B',
             muted='8B8A84', line='DAD4C8', bg=PAPER, bg2='EFEBE3', yellow='E2B866', panel='F2EFE8',
             panel_line='DDD6C8', font_display='Palatino Linotype', page_numbers=False)
    d.page = types.MethodType(ink_page, d)        # skill 版式也用水墨页眉页脚
    d.set_sections(SEC, en=SEC_EN)
    cover(d)
    agenda(d)
    section(d, 0, ['长距离输水工程的规模与风险', '水锤事故的典型危害', '泄漏检测的现实难点'])
    page_background(d)
    page_moc(d)
    page_result(d)
    page_innov(d)
    ending(d)
    path = d.save(OUT, render='--no-render' not in sys.argv)
    print('saved:', path)
    todo = [k for k, v in IMAGES.items() if not v]
    if todo:
        print('占位图（待 imagegen 替换，放入 img/）：', ', '.join(todo))


if __name__ == '__main__':
    build()
