"""Deck: owns the Presentation, theme, sections/navigation state and page chrome.

    from pptstudio import Deck
    d = Deck('academic_blue', scenario='research', nav='top')
    d.meta(title='...', author='...', org='...', date='2026.10', logo='logo.png')
    d.set_sections(['研究背景', '研究方法', '结果分析', '结论展望'])
    d.cover(title='...', subtitle='...', en='...')
    d.agenda()
    d.section(1)
    d.cards('页面结论式标题', items=[{'icon': '水', 'title': '...', 'text': '...'}])
    cv = d.page('自定义页标题')      # blank content page with chrome -> draw with cv.* / components
    d.save('out.pptx')
"""
from __future__ import annotations

import os

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from . import components as K
from . import oxml as X
from . import theme as TH
from .canvas import Canvas, Box
from .units import emu, SLIDE_W, SLIDE_H


class Deck:
    def __init__(self, theme='nsfc_blue', scenario='research', nav=None, page_numbers=True, portable=False,
                 skin=None, **overrides):
        base_style = TH.PRESETS.get(theme, {}).get('style', 'modern') if isinstance(theme, str) else \
            getattr(theme, 'style', 'modern')
        base_style = overrides.get('style', base_style)
        ov = dict(TH.SCENARIO_DEFAULTS.get(scenario, {})) if base_style != 'defense' else {}
        ov.update(overrides)
        if portable:
            ov['portable'] = True
        self.theme = TH.get(theme, **ov)
        if skin:   # 答辩风格家族: chrome of a template family (see theme.SKINS), colors stay the theme's
            self.theme = TH.apply_skin(self.theme, skin)
        self.scenario = scenario
        self.prs = Presentation()
        self.prs.slide_width = Emu(emu(SLIDE_W))
        self.prs.slide_height = Emu(emu(SLIDE_H))
        self._blank = self.prs.slide_layouts[6]
        self._apply_theme_xml()
        self._meta = dict(title=None, subtitle=None, en=None, author=None, advisor=None, org=None, date=None,
                          logo=None, event=None)
        self.sections = []
        self.section_en = []
        self.cur_section = -1
        self.nav = nav
        self.page_numbers = page_numbers
        self.pages = []  # (kind, title)
        self._sub_counter = {}  # section index -> number of content pages (auto '2.3' numbering)
        self._next_num = None   # num= passed to a layout that has no num parameter

    @property
    def style(self):
        return self.theme.style

    # ------------------------------------------------------------ setup
    @property
    def meta(self):
        return _MetaProxy(self)

    def set_meta(self, **kw):
        self._meta.update({k: v for k, v in kw.items()})
        return self

    def add_colors(self, **named):
        """Register named colors usable everywhere a color spec is accepted (incl. '{{name|text}}' markup and
        'name+30' tints), e.g. d.add_colors(moc='0072BD', wcm='77AC30') to keep a paper's method colors."""
        self.theme.extras = dict(self.theme.extras or {})
        self.theme.extras.setdefault('colors', {}).update(named)
        return self

    def set_sections(self, titles, en=None, short=None):
        """Chapter names (agenda, section pages, numbering, navigation). `short`: compact labels for nav='mini'."""
        self.sections = list(titles)
        self.section_en = list(en) if en else [None] * len(titles)
        self.section_short = list(short) if short else [_short_label(x) for x in self.sections]
        return self

    def _apply_theme_xml(self):
        """Write palette & fonts into the theme part so charts, new text and PowerPoint's UI inherit them."""
        t = self.theme
        master = self.prs.slide_master
        theme_part = None
        for rel in master.part.rels.values():
            if rel.reltype.endswith('/theme'):
                theme_part = rel.target_part
        if theme_part is None:
            return
        root = etree.fromstring(theme_part.blob)
        A = X.NS['a']
        cs = root.find('.//{%s}clrScheme' % A)
        mapping = {'dk1': t.dark if not t.is_dark else '000000', 'lt1': 'FFFFFF', 'dk2': t.c('primary_dark'),
                   'lt2': t.bg2 if not t.is_dark else 'E7ECF5', 'accent1': t.primary, 'accent2': t.secondary,
                   'accent3': t.accent, 'accent4': t.c('primary_light'), 'accent5': t.muted, 'accent6': t.c('primary-40'),
                   'hlink': t.primary, 'folHlink': t.c('primary_dark')}
        if cs is not None:
            cs.set('name', 'pptstudio')
            for child in cs:
                nm = etree.QName(child).localname
                if nm in mapping:
                    for c in list(child):
                        child.remove(c)
                    etree.SubElement(child, '{%s}srgbClr' % A, val=mapping[nm])
        fs = root.find('.//{%s}fontScheme' % A)
        if fs is not None:
            fs.set('name', 'pptstudio')
            for kind, fam in (('majorFont', t.f_title), ('minorFont', t.f_body)):
                el = fs.find('{%s}%s' % (A, kind))
                if el is None:
                    continue
                lat = el.find('{%s}latin' % A)
                ea = el.find('{%s}ea' % A)
                if lat is not None:
                    lat.set('typeface', t.f_latin if kind == 'minorFont' else fam)
                if ea is not None:
                    ea.set('typeface', fam)
                for f in el.findall('{%s}font' % A):
                    if f.get('script') in ('Hans', 'Hant'):
                        f.set('typeface', fam)
        theme_part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

    # ------------------------------------------------------------ slides
    def new(self, kind='content', bg=True, bg_kind=None) -> Canvas:
        slide = self.prs.slides.add_slide(self._blank)
        cv = Canvas(self, slide, kind)
        if bg:
            K.background(cv, bg_kind or kind)
        self.pages.append((kind, None))
        return cv

    def page(self, title, en=None, source=None, bg=True, section=None, nav=True, num=None) -> Canvas:
        """Content page with chrome (background, navigation, title, logo, page number, source).
        Returns a Canvas whose .content Box is the free area for the body.
        num (defense style): section number shown before the title; None = auto ('2.3'), False = none."""
        t = self.theme
        if section is not None:
            self.cur_section = section - 1 if isinstance(section, int) else self.sections.index(section)
        if num is None and self._next_num is not None:
            num, self._next_num = self._next_num, None
        if t.style == 'defense':
            return self._page_defense(title, en, source, bg, nav, num)
        cv = self.new('content', bg=bg)
        mx = t.margin_x
        top = t.title_y
        cv.content = Box(mx, t.content_top, SLIDE_W - 2 * mx, t.content_bottom - t.content_top)
        if nav and self.nav == 'top' and self.sections:
            bottom = K.nav_top(cv, self.sections, self.cur_section)
            top = bottom + 0.22
            cv.content = Box(mx, top + 1.02, SLIDE_W - 2 * mx, t.content_bottom - (top + 1.02))
        elif nav and self.nav == 'left' and self.sections:
            sw = K.nav_left(cv, self.sections, self.cur_section)
            mx = sw + 0.55
            cv.content = Box(mx, t.content_top, SLIDE_W - mx - t.margin_x, t.content_bottom - t.content_top)
        if title:
            en_txt = en
            b = K.page_title(cv, title, en_txt, x=mx, y=top, w=SLIDE_W - mx - t.margin_x -
                             (1.6 if (self._meta.get('logo') and self.nav not in ('top', 'left')) else 0))
            if self._meta.get('logo') and self.nav not in ('top', 'left'):
                K.logo(cv)
            cv.content = Box(cv.content.x, max(cv.content.y, b + 0.28), cv.content.w,
                             t.content_bottom - max(cv.content.y, b + 0.28))
        n = len(self.prs.slides)
        K.footer(cv, n if self.page_numbers else None, source)
        self.pages[-1] = ('content', title)
        return cv

    def _page_defense(self, title, en, source, bg, nav, num):
        from . import academic as A
        t = self.theme
        cv = self.new('content', bg=False)
        if bg:
            A.background(cv, 'content')
        if num is None and self.sections and self.cur_section >= 0:
            k = self._sub_counter.get(self.cur_section, 0) + 1
            self._sub_counter[self.cur_section] = k
            num = f'{self.cur_section + 1}.{k}'
        elif num is False:
            num = None
        top = 0.0
        x0 = 0.55
        if nav and self.nav == 'top' and self.sections:
            top = _nav_top_defense(cv, self.sections, self.cur_section)
        elif nav and self.nav == 'left' and self.sections:
            x0 = _nav_left_defense(cv, self.sections, self.cur_section) + 0.35
        if title:
            if top or x0 > 1:
                y = _title_simple(cv, title, num, x0, top + 0.12)
            else:
                mnav = None
                if nav and self.nav == 'mini' and self.sections and self.cur_section >= 0:
                    mnav = (getattr(self, 'section_short', None) or [x[:4] for x in self.sections], self.cur_section)
                y = A.title_bar(cv, title, num, en=en, nav=mnav)
        else:
            y = top + 0.3
        n = len(self.prs.slides)
        A.bottom_rule(cv, n if self.page_numbers else None)
        nav_h = 0.0
        if nav and self.nav == 'bottom' and self.sections and self.cur_section >= 0:
            # chapter chips centred above the bottom rule (开题·答辩 decks)
            labels = getattr(self, 'section_short', None) or [x[:4] for x in self.sections]
            tw = A.mini_nav_width(cv, labels)
            A.mini_nav(cv, labels, self.cur_section, (SLIDE_W + tw) / 2, SLIDE_H - 0.66)
            nav_h = 0.42
        if source:
            cv.text(x0, SLIDE_H - 0.56 - nav_h, 10.5, 0.26, source, size=9, color='muted', wrap=False, fit=True,
                    min_size=7, name='txt:source')
        cv.content = Box(x0, y, SLIDE_W - x0 - 0.55, (SLIDE_H - 0.56 - nav_h - (0.26 if source else 0)) - y)
        self.pages[-1] = ('content', title)
        return cv

    # ------------------------------------------------------------ output
    def save(self, path, render=False, check=True, pdf=False, width=1600):
        path = os.path.abspath(path)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        cp = self.prs.core_properties          # python-pptx's template carries its author's name: reset it
        cp.title = (self._meta.get('title') or '').replace('\n', '')
        cp.author = self._meta.get('author') or ''
        cp.last_modified_by = ''
        cp.comments = ''
        self.prs.save(path)
        report = None
        if check:
            from . import qa
            report = qa.check(path, self.theme)
            if report['issues']:
                print(qa.format_report(report))
        if render or pdf:
            from . import render as R
            try:
                if render:
                    outdir = os.path.splitext(path)[0] + '_preview'
                    pngs = R.export_png(path, outdir, width=width)
                    sheet = R.contact_sheet(pngs, os.path.join(outdir, 'contact_sheet.jpg'))
                    print('preview:', outdir, '\ncontact sheet:', sheet)
                if pdf:
                    print('pdf:', R.to_pdf(path))
            except Exception as e:  # no PowerPoint / COM unavailable
                print(f'[render skipped] PowerPoint COM not available ({e.__class__.__name__}: {e}). '
                      f'Use LibreOffice: soffice --headless --convert-to pdf "{path}" then view the PDF pages.')
        return path


class _MetaProxy:
    """Allows both d.meta(title=...) and d.meta['title'] / d.meta.get('title')."""

    def __init__(self, deck):
        self._d = deck

    def __call__(self, **kw):
        return self._d.set_meta(**kw)

    def __getitem__(self, k):
        return self._d._meta.get(k)

    def get(self, k, default=None):
        v = self._d._meta.get(k)
        return default if v is None else v


def _attach():
    """Attach layout functions as Deck methods: d.cover(...), d.cards(...), ...
    cover / agenda / section / ending dispatch to the 科研答辩风 versions when theme.style == 'defense'."""
    from . import layouts
    from . import layouts_ac as AC
    defense_versions = {'cover': AC.cover, 'agenda': AC.agenda, 'section': AC.section, 'ending': AC.ending}
    for name in list(layouts.REGISTRY):
        fn = layouts.REGISTRY[name]

        def make(f, nm):
            import inspect
            accepts_num = 'num' in inspect.signature(f).parameters

            def method(self, *a, **kw):
                if self.theme.style == 'defense' and nm in defense_versions:
                    return defense_versions[nm](self, *a, **kw)
                if 'num' in kw and not accepts_num:
                    self._next_num = kw.pop('num')
                return f(self, *a, **kw)
            method.__name__ = f.__name__
            method.__doc__ = ((defense_versions[nm].__doc__ or '') + '\n[modern] ' + (f.__doc__ or '')) \
                if nm in defense_versions else f.__doc__
            return method
        setattr(Deck, name, make(fn, name))


_attach()


def _title_simple(cv, title, num, x0, y):
    """Title used with navigation bars: navy bold, short underline mark."""
    t = cv.t
    txt = (num + '  ' if num else '') + title
    cv.text(x0, y, SLIDE_W - x0 - 0.6, 0.6, txt, size=t.size_title, color='primary_deep', role='title',
            anchor='middle', wrap=False, fit=True, min_size=16, name='txt:title')
    cv.line(x0, y + 0.66, SLIDE_W - 0.55, y + 0.66, 'primary+60', 1.0)
    cv.rect(x0, y + 0.63, 1.0, 0.06, fill='primary', name='deco:titlemark')
    return y + 0.86


def _short_label(title):
    """'研究内容与方案' -> '研究内容', '总结与展望' -> '总结', '立项依据' -> '立项依据' (compact nav chips)."""
    for sep in ('与', '及', '和', '、', '：'):
        head = title.split(sep)[0]
        if sep in title and len(head) >= 2:
            return head[:4]
    return title[:4]


def _nav_top_defense(cv, sections, current):
    """Navy navigation band with section tabs; current tab = white chip with navy text. Returns band bottom y."""
    t = cv.t
    h = 0.62
    cv.rect(0, 0, SLIDE_W, h, grad=[('primary_deep', 1.0), ('primary', 1.0)], grad_angle=0, name='deco:navband')
    x = 0.4
    lg = cv.deck.meta.get('logo')
    if lg:
        K.logo(cv, lg, x=x, y=0.09, h=0.44, align='left')
        x += 1.9
    n = len(sections)
    cw = (SLIDE_W - x - 0.3) / n
    for i, s in enumerate(sections):
        on = i == current
        cx = x + i * cw
        if on:
            cv.rect(cx + 0.04, 0.1, cw - 0.08, h - 0.2, fill='FFFFFF', radius=0.04, name='deco:navon')
        cv.text(cx, 0, cw, h, s, size=13, color='primary_deep' if on else 'FFFFFF', bold=on, align='center',
                anchor='middle', wrap=False, fit=True, min_size=8, alpha=1.0 if on else 0.85, name=f'txt:nav{i}')
    return h


def _nav_left_defense(cv, sections, current, w=2.0):
    """Light sidebar with section list; current = navy pill with arrow. Returns sidebar width."""
    cv.rect(0, 0, w, SLIDE_H, fill='panel', name='deco:sidebar')
    lg = cv.deck.meta.get('logo')
    top = 0.9
    if lg:
        K.logo(cv, lg, x=0.3, y=0.35, h=0.5, align='left')
        top = 1.4
    n = len(sections)
    ih = min(0.62, (SLIDE_H - top - 0.8) / max(1, n))
    for i, s in enumerate(sections):
        y = top + i * (ih + 0.12)
        on = i == current
        if on:
            cv.rect(0.22, y, w - 0.44, ih, fill='primary_deep', radius=0.04, name='deco:navpill')
            cv.shape('triangle', w - 0.22, y + ih / 2 - 0.09, 0.14, 0.18, fill='primary_deep', rot=90,
                     name='deco:navarrow')
        cv.text(0.22, y, w - 0.44, ih, s, size=13, color='FFFFFF' if on else 'text', bold=on, align='center',
                anchor='middle', wrap=False, fit=True, min_size=8, name=f'txt:nav{i}')
    return w
