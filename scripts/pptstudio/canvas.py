"""Canvas: drawing API on one slide. All coordinates in INCHES (slide = 13.333 x 7.5).

Colors: theme roles ('primary', 'text', 'primary_pale', 'primary+30' tint, 'primary-20' shade) or hex.
Text markup: '**bold**', '[[accent emphasis]]' (primary color + bold), paragraphs separated by '\n'.
Shape names: prefix 'bg:' / 'deco:' for decorative shapes (QA ignores overlaps), 'txt:' for text.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Emu

from . import fonts as F
from . import icons as I
from . import oxml as X
from .units import emu, SLIDE_W, SLIDE_H

PRST = {
    'rect': MSO_SHAPE.RECTANGLE, 'roundRect': MSO_SHAPE.ROUNDED_RECTANGLE, 'ellipse': MSO_SHAPE.OVAL,
    'triangle': MSO_SHAPE.ISOSCELES_TRIANGLE, 'rtTriangle': MSO_SHAPE.RIGHT_TRIANGLE,
    'parallelogram': MSO_SHAPE.PARALLELOGRAM, 'trapezoid': MSO_SHAPE.TRAPEZOID, 'hexagon': MSO_SHAPE.HEXAGON,
    'octagon': MSO_SHAPE.OCTAGON, 'diamond': MSO_SHAPE.DIAMOND, 'pentagon': MSO_SHAPE.REGULAR_PENTAGON,
    'chevron': MSO_SHAPE.CHEVRON, 'homePlate': MSO_SHAPE.PENTAGON, 'rightArrow': MSO_SHAPE.RIGHT_ARROW,
    'leftArrow': MSO_SHAPE.LEFT_ARROW, 'upArrow': MSO_SHAPE.UP_ARROW, 'downArrow': MSO_SHAPE.DOWN_ARROW,
    'leftRightArrow': MSO_SHAPE.LEFT_RIGHT_ARROW, 'blockArc': MSO_SHAPE.BLOCK_ARC, 'arc': MSO_SHAPE.ARC,
    'donut': MSO_SHAPE.DONUT, 'pie': MSO_SHAPE.PIE, 'chord': MSO_SHAPE.CHORD, 'wave': MSO_SHAPE.WAVE,
    'doubleWave': MSO_SHAPE.DOUBLE_WAVE, 'round1Rect': MSO_SHAPE.ROUND_1_RECTANGLE,
    'round2SameRect': MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 'round2DiagRect': MSO_SHAPE.ROUND_2_DIAG_RECTANGLE,
    'snip1Rect': MSO_SHAPE.SNIP_1_RECTANGLE, 'snip2DiagRect': MSO_SHAPE.SNIP_2_DIAG_RECTANGLE,
    'plaque': MSO_SHAPE.PLAQUE, 'teardrop': MSO_SHAPE.TEAR, 'star5': MSO_SHAPE.STAR_5_POINT,
    'cloud': MSO_SHAPE.CLOUD, 'flowChartProcess': MSO_SHAPE.FLOWCHART_PROCESS,
    'leftBracket': MSO_SHAPE.LEFT_BRACKET, 'rightBracket': MSO_SHAPE.RIGHT_BRACKET,
    'leftBrace': MSO_SHAPE.LEFT_BRACE, 'rightBrace': MSO_SHAPE.RIGHT_BRACE, 'halfFrame': MSO_SHAPE.HALF_FRAME,
    'corner': MSO_SHAPE.CORNER, 'frame': MSO_SHAPE.FRAME, 'can': MSO_SHAPE.CAN, 'cube': MSO_SHAPE.CUBE,
    'flowChartTerminator': MSO_SHAPE.FLOWCHART_TERMINATOR, 'wedgeRoundRectCallout': MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
}


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float

    @property
    def r(self):
        return self.x + self.w

    @property
    def b(self):
        return self.y + self.h

    @property
    def cx(self):
        return self.x + self.w / 2

    @property
    def cy(self):
        return self.y + self.h / 2

    def inset(self, dx, dy=None):
        dy = dx if dy is None else dy
        return Box(self.x + dx, self.y + dy, self.w - 2 * dx, self.h - 2 * dy)

    def cols(self, n, gap=0.24, weights=None):
        weights = weights or [1] * n
        tot = sum(weights)
        avail = self.w - gap * (n - 1)
        out, x = [], self.x
        for wgt in weights:
            w = avail * wgt / tot
            out.append(Box(x, self.y, w, self.h))
            x += w + gap
        return out

    def rows(self, n, gap=0.24, weights=None):
        weights = weights or [1] * n
        tot = sum(weights)
        avail = self.h - gap * (n - 1)
        out, y = [], self.y
        for wgt in weights:
            h = avail * wgt / tot
            out.append(Box(self.x, y, self.w, h))
            y += h + gap
        return out

    def grid(self, n, cols=None, gap=0.24, row_gap=None):
        cols = cols or (n if n <= 4 else (3 if n in (5, 6, 9) else 4))
        rows = (n + cols - 1) // cols
        rg = gap if row_gap is None else row_gap
        out = []
        for i, rbox in enumerate(self.rows(rows, rg)):
            k = min(cols, n - i * cols)
            cs = rbox.cols(cols, gap)
            if k < cols:  # center the last incomplete row
                off = (cols - k) * (cs[0].w + gap) / 2
                cs = [Box(c.x + off, c.y, c.w, c.h) for c in cs[:k]]
            out += cs[:k]
        return out

    def split(self, frac, gap=0.3, vertical=False):
        if vertical:
            h1 = (self.h - gap) * frac
            return Box(self.x, self.y, self.w, h1), Box(self.x, self.y + h1 + gap, self.w, self.h - h1 - gap)
        w1 = (self.w - gap) * frac
        return Box(self.x, self.y, w1, self.h), Box(self.x + w1 + gap, self.y, self.w - w1 - gap, self.h)



_PUNCT_BREAK = set('、，；：,;。）)」』”')


def _orphan_break(tx, lines, want_last=4):
    """Plain-text index at which to insert a soft line break so the last line gets >= want_last chars.
    Only for CJK text (never splits Latin words). Returns None if not applicable."""
    last = lines[-1].strip()
    body = tx.rstrip()
    if len(lines) < 2 or len(last) > 2 or not body.endswith(last):
        return None
    start_last = len(body) - len(last)
    need = max(2, want_last - len(last))
    prev_len = len(lines[-2].strip())
    best = None
    for k in range(need, min(8, prev_len - 3) + 1):
        idx = start_last - k
        if idx <= 0:
            break
        if body[idx - 1] in _PUNCT_BREAK:
            best = idx
            break
    if best is None:
        idx = start_last - need
        if idx <= 0:
            return None
        best = idx
    a, c = body[best - 1], body[best]
    if not (F.is_cjk(a) or a in _PUNCT_BREAK) or not (F.is_cjk(c)):
        return None
    if c in F.NO_START:
        return None
    return best

_MARK = re.compile(r'(\*\*.+?\*\*|\[\[.+?\]\]|\{\{[^{}|]+\|.+?\}\})')


# PowerPoint breaks lines after '-' / '–' even inside numbers ('4.6–|10.5', 'A-|50'); word joiners are ignored.
# U+2011 (non-breaking hyphen) is the only reliable glue, so ranges and codes are rewritten with it.
_NOBREAK = re.compile('(?<=[0-9A-Za-z.%])[\\-–](?=[0-9])')


def nobreak(s: str) -> str:
    """'4.6–10.5' / 'A-50' -> joined with U+2011 so PowerPoint never splits them across lines."""
    return _NOBREAK.sub('‑', s) if s else s


def parse_markup(s: str):
    """'a **b** [[c]] {{wcm|d}}' -> [('a ', None), ('b','bold'), ('c','emph'), ('d', 'color:wcm')]
    {{color|text}}: bold text in a theme role / extras color name / hex (e.g. method colors kept from the figures)."""
    out = []
    for part in _MARK.split(nobreak(s)):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            out.append((part[2:-2], 'bold'))
        elif part.startswith('[[') and part.endswith(']]'):
            out.append((part[2:-2], 'emph'))
        elif part.startswith('{{') and part.endswith('}}') and '|' in part:
            spec, txt = part[2:-2].split('|', 1)
            out.append((txt, 'color:' + spec.strip()))
        else:
            out.append((part, None))
    return out


def plain(s: str) -> str:
    return ''.join(t for t, _ in parse_markup(s))


class Canvas:
    def __init__(self, deck, slide, kind='content'):
        self.deck = deck
        self.slide = slide
        self.t = deck.theme
        self.kind = kind
        self.W, self.H = SLIDE_W, SLIDE_H
        self.last_size = None
        self.content = Box(self.t.margin_x, self.t.content_top, SLIDE_W - 2 * self.t.margin_x,
                           self.t.content_bottom - self.t.content_top)

    # ------------------------------------------------------------ helpers
    def c(self, spec):
        return self.t.c(spec)

    def _style(self, sh, fill=None, alpha=1.0, line=None, line_w=1.0, line_alpha=1.0, dash=None, grad=None,
               grad_angle=90, grad_kind='linear', shadow=None, glow=None, soft_edge=None, rot=0, name=None):
        X.strip_style(sh)
        if grad:
            stops = []
            for i, g in enumerate(grad):
                if isinstance(g, (tuple, list)):
                    if len(g) == 3:
                        pos, col, al = g
                    else:
                        col, al = g
                        pos = 100 * i / max(1, len(grad) - 1)
                else:
                    col, al, pos = g, 1.0, 100 * i / max(1, len(grad) - 1)
                stops.append((pos, self.c(col), al))
            X.set_grad(sh, stops, grad_angle, grad_kind)
        else:
            X.set_fill(sh, self.c(fill) if fill else None, alpha)
        if line:
            X.set_line(sh, self.c(line), line_w, line_alpha, dash=dash)
        else:
            X.set_line(sh, None)
        X.clear_effects(sh)
        if shadow:
            self._shadow(sh, shadow)
        if glow:
            gc, ga, gr = (glow + (None, None))[:3] if isinstance(glow, tuple) else (glow, 0.45, 12)
            X.set_glow(sh, self.c(gc), ga or 0.45, gr or 12)
        if soft_edge:
            X.set_soft_edge(sh, soft_edge)
        if rot:
            X.set_rotation(sh, rot)
        if name:
            sh.name = name
        return sh

    def _shadow(self, sh, spec):
        t = self.t
        if spec is True or spec == 'soft':
            X.set_shadow(sh, self.c('primary_dark') if not t.is_dark else '000000', t.shadow_alpha, t.shadow_blur,
                         t.shadow_dist, 90)
        elif spec == 'strong':
            X.set_shadow(sh, self.c('primary_dark') if not t.is_dark else '000000', min(0.35, t.shadow_alpha * 1.8),
                         t.shadow_blur * 1.2, t.shadow_dist * 1.6, 90)
        elif spec == 'glow':
            X.set_shadow(sh, self.c('primary'), 0.35, 28, 0, 90)
        elif isinstance(spec, dict):
            X.set_shadow(sh, self.c(spec.get('color', 'primary_dark')), spec.get('alpha', 0.15), spec.get('blur', 20),
                         spec.get('dist', 4), spec.get('dir', 90), spec.get('size', 1.0))

    # ------------------------------------------------------------ shapes
    def shape(self, prst, x, y, w, h, fill='primary', adj=None, name=None, **kw):
        """Any preset geometry (see PRST). adj: dict like {'adj': 50000} or {'adj1':..,'adj2':..}."""
        sh = self.slide.shapes.add_shape(PRST[prst], emu(x), emu(y), emu(max(w, 0.001)), emu(max(h, 0.001)))
        if adj:
            X.set_adj(sh, **adj)
        self._style(sh, fill=fill, name=name, **kw)
        if sh.has_text_frame:
            X.set_body(sh, 0, 0, 0, 0, 'ctr')
        return sh

    def rect(self, x, y, w, h, fill='primary', radius=0.0, name=None, **kw):
        """radius in inches (0 = sharp). Rounded corners keep a constant radius regardless of size."""
        if radius and radius > 0:
            adj = min(50000, int(radius / max(0.001, min(w, h)) * 100000))
            return self.shape('roundRect', x, y, w, h, fill, adj={'adj': adj}, name=name, **kw)
        return self.shape('rect', x, y, w, h, fill, name=name, **kw)

    def ellipse(self, x, y, w, h, fill='primary', name=None, **kw):
        return self.shape('ellipse', x, y, w, h, fill, name=name, **kw)

    def circle(self, cx, cy, r, fill='primary', name=None, **kw):
        return self.shape('ellipse', cx - r, cy - r, 2 * r, 2 * r, fill, name=name, **kw)

    def ring(self, cx, cy, r, thickness, color='primary', start=0, sweep=360, name=None, **kw):
        """Arc band (progress ring) from angle start (deg, 0 = 12 o'clock, clockwise) sweeping `sweep` deg."""
        if sweep >= 359.99:
            adj = int(thickness / (2 * r) * 100000)
            return self.shape('donut', cx - r, cy - r, 2 * r, 2 * r, color, adj={'adj': adj}, name=name, **kw)
        # blockArc angles: 0 = 3 o'clock, clockwise positive (in 60000ths)
        a1 = (start - 90) % 360
        a2 = (start + sweep - 90) % 360
        adj = {'adj1': int(a1 * 60000), 'adj2': int(a2 * 60000), 'adj3': int(thickness / (2 * r) * 100000)}
        return self.shape('blockArc', cx - r, cy - r, 2 * r, 2 * r, color, adj=adj, name=name, **kw)

    def line(self, x1, y1, x2, y2, color='line', width=1.0, alpha=1.0, dash=None, head=None, tail=None,
             cap='round', name=None, grad=None):
        cn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        X.strip_style(cn)
        stops = None
        if grad:
            stops = [(100 * i / max(1, len(grad) - 1), self.c(g[0] if isinstance(g, tuple) else g),
                      g[1] if isinstance(g, tuple) else 1.0) for i, g in enumerate(grad)]
        X.set_line(cn, self.c(color), width, alpha, dash=dash, cap=cap, head=head, tail=tail, grad_stops=stops,
                   grad_angle=0 if abs(x2 - x1) >= abs(y2 - y1) else 90)
        if name:
            cn.name = name
        return cn

    def path(self, subpaths, x, y, w, h, vb_w, vb_h, fill=None, line=None, line_w=1.0, alpha=1.0,
             line_alpha=1.0, name=None, filled_idx=None, cap='round', **kw):
        """Custom geometry from subpaths in a vb_w x vb_h coordinate space."""
        sh = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(max(w, 0.001)), emu(max(h, 0.001)))
        X.strip_style(sh)
        X.set_custgeom(sh, subpaths, vb_w, vb_h, fill_mode='norm' if (fill or kw.get('grad')) else 'none',
                       filled_idx=filled_idx)
        self._style(sh, fill=fill, alpha=alpha, name=name, **{k: v for k, v in kw.items() if k in (
            'grad', 'grad_angle', 'grad_kind', 'shadow', 'glow', 'soft_edge', 'rot')})
        if line:
            X.set_line(sh, self.c(line), line_w, line_alpha, cap=cap)
        return sh

    def poly(self, pts, fill='primary', line=None, line_w=1.0, closed=True, smooth=False, name=None, **kw):
        """Polygon/polyline through points (inches, absolute). smooth=True -> Catmull-Rom curve."""
        xs, ys = [p[0] for p in pts], [p[1] for p in pts]
        x0, y0 = min(xs), min(ys)
        w, h = max(max(xs) - x0, 0.01), max(max(ys) - y0, 0.01)
        S = 1000.0
        rel = [((px - x0) / w * S, (py - y0) / h * S) for px, py in pts]
        segs = []
        if smooth and len(rel) > 2:
            P = [rel[0]] + rel + [rel[-1]]
            for i in range(1, len(P) - 2):
                p0, p1, p2, p3 = P[i - 1], P[i], P[i + 1], P[i + 2]
                c1 = (p1[0] + (p2[0] - p0[0]) / 6, p1[1] + (p2[1] - p0[1]) / 6)
                c2 = (p2[0] - (p3[0] - p1[0]) / 6, p2[1] - (p3[1] - p1[1]) / 6)
                segs.append(('C', c1[0], c1[1], c2[0], c2[1], p2[0], p2[1]))
        else:
            segs = [('L', px, py) for px, py in rel[1:]]
        sub = [{'start': rel[0], 'segs': segs, 'closed': closed}]
        if not closed:
            kw.pop('grad', None)
        return self.path(sub, x0, y0, w, h, S, S, fill=fill if closed else None, line=line, line_w=line_w,
                         name=name, **kw)

    # ------------------------------------------------------------ text
    def _font_for(self, role):
        return self.t.font(role)

    def text(self, x, y, w, h, content, size=None, color='text', role='body', bold=None, italic=False,
             align='left', anchor='top', ls=None, sa=None, sb=None, fit=True, min_size=None, max_lines=None,
             spacing=None, emph=None, name=None, alpha=1.0, grad=None, grad_angle=0, bullet=None,
             bullet_color='primary', indent=None, shadow=None, wrap=True, outline=None, vert=None, rot=0,
             insets=(0, 0, 0, 0), bullet_font=None):
        """Add a text box. content: str (\\n = new paragraph) | list[str|dict]. Returns shape.
        size: pt (default theme body). fit: shrink font so text fits the box (never below min_size).
        dict paragraph keys: text, size, color, bold, role, align, bullet, sa (space after pt)."""
        t = self.t
        size = size or t.size_body
        emph = emph or 'emph'
        paras = content if isinstance(content, (list, tuple)) else str(content).split('\n')
        paras = [p if isinstance(p, dict) else {'text': str(p)} for p in paras]
        ea, lat, force_bold = self._font_for(role)
        b = force_bold if bold is None else bold
        ls = ls if ls is not None else (t.line_spacing if len(paras) > 1 or len(plain(paras[0]['text'])) > 24 else 1.0)
        sa = sa if sa is not None else (0.35 * size if len(paras) > 1 else 0)
        # --- fit
        final = size
        breaks = {}  # paragraph index -> plain-text index where a soft line break is inserted (orphan control)
        if fit and wrap:
            min_size = min_size if min_size is not None else min(size, t.size_min)
            wpt = (w - (insets[0] + insets[2]) / 72) * 72 * 0.975
            if bullet:
                wpt -= size * 1.1
            hpt = (h - (insets[1] + insets[3]) / 72) * 72
            texts = [plain(p['text']) for p in paras]
            rel = [p.get('size', size) / size for p in paras]
            s = size
            while True:
                tot, lines = 0.0, 0
                for i, (tx, r) in enumerate(zip(texts, rel)):
                    ps = s * r
                    hh, n, _ = F.layout_height([tx], wpt, ps, ea, lat, b or paras[i].get('bold', False), ls)
                    tot += hh + (sa * s / size if i < len(texts) - 1 else 0)
                    lines += n
                if (tot <= hpt + 0.5 and (max_lines is None or lines <= max_lines)) or s <= min_size:
                    break
                s -= 0.5
            # avoid orphans (last line of <= 2 chars): shrink at most ~5 %, otherwise break the line earlier
            def orphans(sz):
                out = {}
                for i, (tx, r) in enumerate(zip(texts, rel)):
                    lns = F.wrap(tx, wpt, sz * r, ea, lat, b or paras[i].get('bold', False))
                    if len(lns) > 1 and len(lns[-1].strip()) <= 2:
                        out[i] = lns
                return out
            orph = orphans(s)
            s2 = s
            while orph and s2 - 0.5 >= max(min_size, s * 0.95):
                s2 -= 0.5
                o2 = orphans(s2)
                if not o2:
                    s, orph = s2, {}
                    break
            for i, lns in orph.items():
                k = _orphan_break(texts[i], lns)
                if k is not None:
                    breaks[i] = k
            final = s
        elif fit and not wrap:
            min_size = min_size if min_size is not None else min(size, t.size_min)
            wpt = w * 72 * 0.98
            s = size
            while s > min_size and max(F.text_width(plain(p['text']), s * p.get('size', size) / size, ea, lat, b,
                                                     spacing or 0) for p in paras) > wpt:
                s -= 0.5
            final = s
        scale = final / size
        self.last_size = final
        tb = self.slide.shapes.add_textbox(emu(x), emu(y), emu(max(w, 0.05)), emu(max(h, 0.05)))
        X.set_body(tb, *insets, anchor={'top': 't', 'middle': 'ctr', 'center': 'ctr', 'bottom': 'b', 't': 't',
                                          'ctr': 'ctr', 'b': 'b'}[anchor], wrap=wrap, vert=vert)
        tf = tb.text_frame
        for i, pd in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            psize = pd.get('size', size) * scale
            X.set_para(p, align=pd.get('align', align), line_spacing=pd.get('ls', ls),
                       space_after=(pd.get('sa', sa) * scale if i < len(paras) - 1 else 0),
                       space_before=(sb if sb is not None else None),
                       bullet=(pd.get('bullet', bullet) if pd.get('bullet', bullet) is not None else ''),
                       bullet_color=self.c(bullet_color) if (pd.get('bullet', bullet)) else None,
                       bullet_font=bullet_font,
                       margin_left=(psize * 1.05 if pd.get('bullet', bullet) else (indent or 0)),
                       indent=(-psize * 1.05 if pd.get('bullet', bullet) else 0))
            pr_ea, pr_lat, pr_fb = self._font_for(pd.get('role', role))
            pbold = pd.get('bold', pr_fb if bold is None else bold)
            pcolor = self.c(pd.get('color', color))
            brk = breaks.get(i)
            offset = 0
            for seg, kind in parse_markup(pd['text']):
                pieces = [seg]
                if brk is not None and offset < brk < offset + len(seg):
                    pieces = [seg[:brk - offset], None, seg[brk - offset:]]
                elif brk is not None and brk == offset and offset > 0:
                    pieces = [None, seg]
                offset += len(seg)
                col = pcolor
                rb = pbold
                if kind == 'bold':
                    rb = True
                elif kind == 'emph':
                    rb = True
                    col = self.c(emph)
                elif kind and kind.startswith('color:'):
                    rb = True
                    col = self.c(kind[6:])
                for piece in pieces:
                    if piece is None:
                        p.add_line_break()
                        continue
                    if not piece:
                        continue
                    r = p.add_run()
                    r.text = piece
                    X.style_run(r, size=psize, bold=rb, italic=italic, hexclr=None if grad else col, alpha=alpha,
                                ea=pr_ea, latin=pr_lat, spacing=spacing,
                                grad_stops=[(100 * k / max(1, len(grad) - 1), self.c(g), 1.0)
                                            for k, g in enumerate(grad)] if grad else None, grad_angle=grad_angle,
                                outline=(self.c(outline[0]), outline[1]) if outline else None)
            if not pd['text']:
                X.set_end_para_font(p, psize, pr_ea, pr_lat)
        if shadow:
            self._shadow(tb, shadow)
        if rot:
            X.set_rotation(tb, rot)
        tb.name = name or 'txt:' + plain(paras[0]['text'])[:18]
        return tb

    def measure(self, content, w, size=None, role='body', bold=None, ls=None, sa=None):
        """Height (inches) needed to render content in width w at the given size (no shrinking)."""
        t = self.t
        size = size or t.size_body
        paras = content if isinstance(content, (list, tuple)) else str(content).split('\n')
        paras = [p['text'] if isinstance(p, dict) else str(p) for p in paras]
        ea, lat, fb = self._font_for(role)
        b = fb if bold is None else bold
        ls = ls if ls is not None else (t.line_spacing if len(paras) > 1 or len(plain(paras[0])) > 24 else 1.0)
        sa = sa if sa is not None else (0.35 * size if len(paras) > 1 else 0)
        tot = 0
        for i, tx in enumerate(paras):
            hh, n, _ = F.layout_height([plain(tx)], w * 72 * 0.975, size, ea, lat, b, ls)
            tot += hh + (sa if i < len(paras) - 1 else 0)
        return tot / 72

    def width_of(self, s, size, role='body', bold=None, spacing=0):
        ea, lat, fb = self._font_for(role)
        return F.text_width(plain(s), size, ea, lat, fb if bold is None else bold, spacing) / 72

    # ------------------------------------------------------------ icons
    def icon(self, name, x, y, size, color='primary', stroke=2.0, alpha=1.0, name_prefix='icon:'):
        """Lucide icon as native vector (square box at x,y with side `size` inches). name can be a keyword."""
        nm = I.resolve(name)
        subs, filled = I.subpaths(nm)
        sh = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(size), emu(size))
        X.strip_style(sh)
        X.set_custgeom(sh, subs, 24, 24, fill_mode='none', filled_idx=set(filled))
        col = self.c(color)
        X.set_fill(sh, col, alpha)
        X.set_line(sh, col, max(0.5, size * 72 * stroke / 24), alpha, cap='round', join='round')
        X.clear_effects(sh)
        sh.name = name_prefix + nm
        return sh

    # ------------------------------------------------------------ images
    def image(self, path, x, y, w, h, fit='cover', focus=(0.5, 0.5), radius=0.0, shape=None, alpha=1.0,
              line=None, line_w=1.0, shadow=None, name=None):
        """Insert picture. fit: cover (crop to fill box) | contain (letterbox inside box) | stretch.
        shape: None | 'ellipse' | 'roundRect' | any preset name; radius (in) for rounded corners."""
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        bx, by, bw, bh = x, y, w, h
        if fit == 'contain':
            s = min(w / iw, h / ih)
            bw, bh = iw * s, ih * s
            bx, by = x + (w - bw) / 2, y + (h - bh) / 2
        pic = self.slide.shapes.add_picture(path, emu(bx), emu(by), emu(bw), emu(bh))
        if fit == 'cover':
            ia, ba = iw / ih, w / h
            if ia > ba:  # too wide -> crop left/right
                keep = ba / ia
                off = max(0.0, min(1 - keep, focus[0] - keep / 2))
                pic.crop_left, pic.crop_right = off, 1 - keep - off
            elif ia < ba:
                keep = ia / ba
                off = max(0.0, min(1 - keep, focus[1] - keep / 2))
                pic.crop_top, pic.crop_bottom = off, 1 - keep - off
        if shape or radius:
            if radius and not shape:
                X.pic_geom(pic, 'roundRect', adj=min(50000, int(radius / max(0.01, min(bw, bh)) * 100000)))
            else:
                X.pic_geom(pic, shape)
        if alpha < 1:
            X.pic_alpha(pic, alpha)
        if line:
            X.set_line(pic, self.c(line), line_w)
        if shadow:
            self._shadow(pic, shadow)
        pic.name = name or 'img:' + os.path.basename(path)[:20]
        return pic

    def equation(self, tex, x, y, size=20.0, color='text', w=None, h=None, align='left', anchor='top', name=None):
        """Typeset a LaTeX-style formula (mathtext) as a picture at true point size; shrinks to fit w/h if given.
        Returns the Box actually covered by the formula."""
        from . import mathtex
        path, ew, eh = mathtex.render(tex, size, self.c(color))
        s = 1.0
        if w:
            s = min(s, w / ew)
        if h:
            s = min(s, h / eh)
        ew, eh = ew * s, eh * s
        bx = x if not w else {'left': x, 'center': x + (w - ew) / 2, 'right': x + w - ew}[align]
        by = y if not h else {'top': y, 'middle': y + (h - eh) / 2, 'bottom': y + h - eh}[anchor]
        pic = self.slide.shapes.add_picture(path, emu(bx), emu(by), emu(ew), emu(eh))
        pic.name = name or 'eq:' + re.sub(r'[^A-Za-z0-9]', '', tex)[:16]
        return Box(bx, by, ew, eh)

    # ------------------------------------------------------------ misc
    def notes(self, text):
        self.slide.notes_slide.notes_text_frame.text = text

    def transition(self, kind='fade'):
        X.set_transition(self.slide, kind)

    def to_back(self, shape):
        tree = self.slide.shapes._spTree
        tree.remove(shape._element)
        tree.insert(2, shape._element)

    def to_front(self, shape):
        tree = self.slide.shapes._spTree
        tree.remove(shape._element)
        tree.append(shape._element)
