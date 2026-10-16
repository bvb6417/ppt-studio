"""Galleries: (1) layout catalog — one sample page per layout; (2) style showcase — the same 9-page sample talk
rendered in every 答辩风格家族 (skin). Sample content is a fictional project with synthetic figures
(`pptstudio.samples`), so the thumbnails look like real NSFC / 科技奖 decks instead of grey wireframes."""
from __future__ import annotations

import os

from .deck import Deck

TITLE = '深埋长隧洞围岩灾变机理\n与智能预警关键技术'
NOTE = '注：示例项目与示例数据，仅用于版式演示'


def _meta(d, project='国家自然科学基金面上项目 · 结题汇报（示例）'):
    d.meta(title=TITLE, author='张三', unit='示例大学 水利水电学院', lab='示例省 岩土工程重点实验室',
           project=project, completers='张三、李四、王五、赵六、钱七',
           units=['示例大学', '示例工程研究院', '示例水利水电工程局', '示例省重点实验室'], date='2026年10月')
    d.set_sections(['立项依据', '研究内容与方案', '研究成果与创新', '总结与展望'])


def _figs():
    from . import samples
    return samples.all_figures()


# ------------------------------------------------------------------ shared sample pages
def _p_need(d, F, **kw):
    return d.panels('国家重大需求', statement='深埋长隧洞建设规模空前，[[岩爆]]与[[大变形]]已成为制约工程安全的首要难题',
                    panels=[{'header': '深埋长隧洞建设规模空前', 'image': F['route_map'], 'caption': '图1 示例隧洞线路与地形',
                             'points': ['典型工程隧洞长度 [[30 km]] 以上', '最大埋深超过 [[2000 m]]，地应力高']},
                            {'header': '围岩灾害事件逐年增加',
                             'chart': {'kind': 'column', 'categories': ['2019', '2020', '2021', '2022', '2023', '2024'],
                                       'series': {'岩爆': [12, 15, 19, 23, 27, 31], '大变形': [8, 9, 13, 14, 18, 20]},
                                       'number_format': '0', 'legend': True, 'colors': ['primary', 'emph'],
                                       'gradient': False, 'gap': 80},
                             'image_h': 0.64, 'points': ['近五年灾害事件年均增长 [[21%]]', '造成工期延误与重大安全风险']}],
                    conclusion='亟需揭示深埋围岩灾变机理，建立智能预警理论与技术体系', conclusion_style='yellow',
                    source=NOTE, **kw)


def _p_problems(d, F, **kw):
    return d.problems('关键科学问题', lead='项目拟解决两大[[关键科学问题]]', items=[
        {'label': '问题一', 'question': '高地应力下围岩破裂如何孕育、扩展并诱发岩爆？', 'image': F['stress_field'],
         'caption': '图2 开挖后切向应力集中', 'tag': '研究目标', 'points': ['揭示多尺度破裂演化机理', '建立[[能量释放]]判据']},
        {'label': '问题二', 'question': '如何从多源监测信息中实时识别灾变前兆？', 'image': F['microseismic'],
         'caption': '图3 微震波形与时频特征', 'tag': '研究目标', 'points': ['提取灾变[[前兆特征]]', '实现分钟级预警']}], **kw)


def _p_route(d, F, **kw):
    return d.route('技术路线', target='[[科学目标：]]揭示深埋围岩灾变机理，构建“感知—诊断—预警”一体化方法',
                   top='深埋长隧洞围岩灾变智能预警平台与工程示范', labels=('机理驱动', '数据驱动'),
                   questions=[{'title': '关键科学问题二', 'focus': '灾变前兆实时识别'},
                              {'title': '关键科学问题一', 'focus': '围岩破裂演化机理'}],
                   topics=[['课题4：多源信息融合的前兆识别[[（诊断）]]', '课题3：分布式智能感知网络[[（感知）]]'],
                           ['课题2：岩爆能量释放判据[[（机理）]]', '课题1：多尺度破裂演化试验[[（试验）]]']],
                   left_image=F['section'], right_image=F['network'], left_caption='监测断面布置',
                   right_caption='智能感知网络', **kw)


def _p_method(d, F, **kw):
    return d.figure_text('围岩灾变风险评估模型', image=F['risk3d'], caption='图4 岩爆风险指数响应面',
                         statement='综合[[埋深]]、[[侧压系数]]与岩体完整性，实现隧洞分段风险定量评估', blocks=[
                             {'tag': '多因素\n耦合', 'title': '风险响应面', 'text': '埋深与侧压系数共同控制风险，埋深大于 1500 m 时[[显著升高]]'},
                             {'tag': '分级\n预警', 'title': '三级预警阈值', 'text': '按风险指数划分蓝、黄、红三级，联动处置预案'},
                             {'tag': '工程\n验证', 'title': '示范段验证', 'text': '风险分区与实际灾害位置吻合率 [[87%]]'}],
                         source=NOTE, **kw)


def _p_results(d, F, **kw):
    return d.panels('预警效果验证', statement='本文模型精度显著优于经验模型，典型断面可[[提前 6 d]]发出位移预警',
                    panels=[{'header': '位移监测与分级预警', 'image': F['curves'], 'caption': '图5 典型断面位移曲线',
                             'points': ['红、黄两级阈值联动处置', '预警时刻早于失稳 [[6 d]]']},
                            {'header': '预测精度对比',
                             'table': {'columns': ['模型', 'R²', '平均误差', '预警提前量'],
                                       'rows': [['经验模型', '0.71', '18.6%', '1.5 d'], ['机器学习', '0.88', '9.4%', '3.2 d'],
                                                ['本文模型', '[[0.96]]', '[[4.1%]]', '[[6.0 d]]']], 'h': 0.5},
                             'points': ['误差降低至经验模型的 [[1/4]]', '示范段 [[12 次]]灾害全部提前预警']}],
                    conclusion='“机理+数据”双驱动预警，实现从被动处置到主动防控', source=NOTE, **kw)


def _p_innov(d, F, **kw):
    return d.innovations('特色与创新', subtitle='从“经验判识”到“机理+数据”双驱动的智能预警', items=[
        {'tag': '理论创新', 'text': '建立深埋围岩多尺度破裂演化模型与[[能量释放]]判据，揭示岩爆孕育机理'},
        {'tag': '方法创新', 'text': '提出多源监测信息融合的灾变前兆识别方法，预警提前量提高至 [[6 d]]'},
        {'tag': '技术创新', 'text': '研发分布式智能感知网络与预警云平台，实现分钟级[[闭环预警]]'}],
        footer='为深埋长隧洞安全建设与运行提供[[理论支撑]]与技术保障', image=F['scatter'], caption='图6 预测与实测对比',
        **kw)


# ------------------------------------------------------------------ (1) layout catalog
def _defense(d):
    F = _figs()
    _meta(d)
    d.cover(notes='[gallery] cover(variant="band")')
    d.cover(variant='sky', notes='[gallery] cover(variant="sky")')
    d.cover(variant='award', title='示例省科学技术进步奖', subtitle='深埋长隧洞围岩灾变智能预警关键技术及应用',
            notes='[gallery] cover(variant="award") 科技奖')
    d.cover(variant='left', notes='[gallery] cover(variant="left") 亮蓝白')
    d.cover(variant='tech', project='国家重点研发计划 · 项目答辩（示例）', notes='[gallery] cover(variant="tech") 重点研发')
    d.cover(motif='wave', project='水利类项目 · 结题汇报（示例）', notes='[gallery] cover(motif="wave") 水利/流域主题')
    d.cover(variant='plaque', project='2026年度示例省科学技术进步奖答辩汇报', title='示例省科学技术进步奖',
            subtitle='深埋长隧洞围岩灾变智能预警关键技术及应用', notes='[gallery] cover(variant="plaque") 科技进步奖')
    d.cover(variant='midband', notes='[gallery] cover(variant="midband") 科学技术奖 / 万人计划')
    d.cover(variant='arc', project='国家重点研发计划 · 项目答辩（示例）', notes='[gallery] cover(variant="arc") 重点研发')
    d.cover(variant='tag', project='深埋工程灾变预警方向（示例）', title='青年拔尖人才项目答辩', subtitle='候选人汇报（示例）',
            notes='[gallery] cover(variant="tag") 长江学者 / 人才计划')
    d.cover(variant='lab', project='2026年度示例省重点实验室评估答辩', title='重点实验室建设运行方案',
            subtitle='深部岩体灾变机理与智能防控（示例）', presenter='实验室主任：张三',
            notes='[gallery] cover(variant="lab") 重点实验室')
    d.agenda(descs=['工程需求与科学问题', '技术路线与关键方法', '主要成果与创新点', '结论与应用前景'],
             notes='[gallery] agenda(variant="vertical")')
    d.agenda(variant='split', descs=['Background', 'Methods', 'Results', 'Outlook'], notes='[gallery] agenda(variant="split")')
    d.agenda(variant='card', current=1, notes='[gallery] agenda(variant="card", current=1)')
    d.agenda(variant='corner', notes='[gallery] agenda(variant="corner") 重点研发 汇报目录')
    d.section(1, notes='[gallery] section(i) = agenda with current chapter in red')
    _p_need(d, F, notes='[gallery] panels(image + native chart + conclusion)')
    d.panels('六大研究方向', statement='围绕[[“机理—感知—预警”]]主线形成六个研究方向', stmt_style='centered', cols=3, panels=[
        {'header': h, 'icon': ic, 'points': pts} for h, ic, pts in (
            ('破裂演化机理', 'activity', ['真三轴试验与声发射', '多尺度[[损伤演化]]']),
            ('能量释放判据', 'zap', ['弹性应变能评价', '岩爆[[倾向性]]分级']),
            ('分布式感知', 'radar', ['光纤 + 微震联合', '[[全断面]]覆盖']),
            ('前兆识别', 'scan-search', ['时频特征提取', '[[深度学习]]识别']),
            ('风险评估', 'gauge', ['多因素响应面', '分段[[量化]]评估']),
            ('预警平台', 'cloud', ['云边协同计算', '[[分钟级]]推送']))],
        conclusion='六个方向相互支撑，形成完整技术链条', notes='[gallery] panels(cols=3, icons)')
    d.columns('卡脖子问题', statement='首当其冲的是[[“看不见、判不准、报不及”]]：灾变前兆难以及时捕捉', cols=[
        {'header': '感知范围有限', 'subtitle': '“看不见”', 'points': ['点式传感器覆盖不足', '深部[[信息盲区]]大']},
        {'header': '机理认识不清', 'subtitle': '“判不准”', 'points': ['破裂演化过程复杂', '判据依赖[[经验]]']},
        {'header': '数据孤岛', 'subtitle': '“融不了”', 'points': ['多源数据格式各异', '缺少[[统一平台]]']},
        {'header': '预警滞后', 'subtitle': '“报不及”', 'points': ['人工分析耗时', '处置窗口[[不足 1 d]]']}],
        bottom=[{'tag': '感知\n技术', 'title': '瓶颈一：深部围岩状态难以连续感知', 'text': '需发展分布式、长距离、低功耗监测技术'},
                {'tag': '预警\n模型', 'title': '瓶颈二：前兆识别依赖人工经验', 'text': '需建立机理与数据融合的智能识别模型'}],
        notes='[gallery] columns')
    _p_problems(d, F, notes='[gallery] problems(rows)')
    d.problems('三大研究重点', lead='围绕三个[[关键环节]]展开研究', items=[
        {'label': '重点一', 'question': '破裂演化如何量化？', 'image': F['stress_field'], 'caption': '应力集中', 'tag': '研究重点',
         'points': ['真三轴试验', '[[能量]]演化']},
        {'label': '重点二', 'question': '前兆信号如何识别？', 'image': F['microseismic'], 'caption': '微震时频', 'tag': '研究重点',
         'points': ['时频特征', '[[智能]]识别']},
        {'label': '重点三', 'question': '预警效果如何验证？', 'image': F['curves'], 'caption': '位移预警', 'tag': '研究重点',
         'points': ['示范段验证', '[[闭环]]反馈']}], notes='[gallery] problems(columns)')
    _p_method(d, F, notes='[gallery] figure_text(image + blocks)')
    _p_route(d, F, notes='[gallery] route (with side images)')
    _p_results(d, F, notes='[gallery] panels(image + table)')
    _p_innov(d, F, notes='[gallery] innovations(+ image)')
    d.achievements('代表性成果', headline='在《Tunnelling and Underground Space Technology》等期刊发表系列论文（示例）',
                   images=[F['stress_field'], F['microseismic'], F['risk3d'], F['network']],
                   captions=['破裂演化机理', '前兆识别方法', '风险评估模型', '预警平台'],
                   stats=[{'value': '36', 'unit': '篇', 'label': 'SCI/EI论文'}, {'value': '1200', 'unit': '次', 'label': '他引'},
                          {'value': '18', 'unit': '项', 'label': '授权发明专利'}, {'value': '3', 'unit': '项', 'label': '省部级科技奖'}],
                   notes='[gallery] achievements')
    d.schedule('研究计划与进度安排', periods=['2026', '2027', '2028'], sub_periods=['上半年', '下半年'], tasks=[
        {'name': '研究内容一', 'desc': '破裂演化机理', 'span': (0, 3), 'label': '真三轴试验与数值模拟'},
        {'name': '研究内容二', 'desc': '前兆识别方法', 'span': (2, 5), 'label': '多源融合与模型训练'},
        {'name': '研究内容三', 'desc': '预警平台示范', 'span': (4, 6), 'label': '平台开发与工程验证'}],
        statement='三项研究内容[[循序推进、相互衔接]]，2028 年完成工程示范', notes='[gallery] schedule')
    d.profile('申请人基本情况', '张 三', role='博士研究生（在读）', direction='岩石力学、隧洞灾害预警、智能监测',
              groups=[{'label': '教育经历', 'rows': [('2023-至今', '示例大学', '岩土工程', '博士'),
                                                     ('2020-2023', '示例大学', '水利工程', '硕士')]},
                      {'label': '成果列表', 'rows': [('学术论文', '发表 SCI 论文 4 篇', '', '第一作者'),
                                                     ('技术专利', '授权发明专利 3 项', '', '第一发明人'),
                                                     ('学科竞赛', '全国研究生创新大赛', '', '一等奖')]}],
              conclusion='研究聚焦深埋工程安全这一国家重大需求中的关键科学问题', notes='[gallery] profile')
    d.layers('智能预警系统架构', statement='以[[“感知—传输—诊断—决策”]]为主线构建四层体系', levels=[
        {'label': '决策层', 'items': ['分级预警', '处置预案', '辅助决策']}, {'label': '诊断层', 'items': ['前兆识别', '风险评估']},
        {'label': '传输层', 'items': ['边缘网关', '5G 回传', '云平台']}, {'label': '感知层', 'items': ['微震监测', '光纤应变']}],
        side={'header': '关键技术', 'points': ['多源数据融合', '深度学习识别', '云边协同计算']}, notes='[gallery] layers')
    d.milestones('项目进展', statement='三年来按计划完成全部研究任务，[[多项指标]]超额完成', stats=[
        {'value': '32.6', 'unit': 'km', 'label': '示范段监测里程'}, {'value': '12', 'unit': '次', 'label': '成功预警灾害'},
        {'value': '6', 'unit': 'd', 'label': '平均预警提前量'}],
        events=[{'date': '2023.06', 'title': '项目启动', 'text': '完成方案论证'}, {'date': '2024.03', 'title': '试验完成', 'text': '真三轴试验 120 组'},
                {'date': '2025.05', 'title': '平台上线', 'text': '示范段部署'}, {'date': '2026.06', 'title': '通过验收', 'text': '指标全部达成'}],
        highlight=3, notes='[gallery] milestones')
    d.table_page('预测精度对比', ['模型', 'R²', '平均误差', '最大误差', '预警提前量'],
                 [['经验公式法', '0.71', '18.6%', '42.3%', '1.5 d'], ['支持向量机', '0.84', '11.2%', '27.5%', '2.6 d'],
                  ['长短期记忆网络', '0.88', '9.4%', '21.8%', '3.2 d'], ['本文模型', '[[0.96]]', '[[4.1%]]', '[[9.7%]]', '[[6.0 d]]']],
                 style='three_line', highlight_col=4, note=NOTE, notes='[gallery] table_page(three_line, highlight_col)')
    d.chart_page('灾害类型统计', {'kind': 'column', 'categories': ['轻微岩爆', '中等岩爆', '强烈岩爆', '大变形', '塌方'],
                                'series': {'次数': [46, 21, 7, 18, 9]}, 'highlight': 2},
                 points=['强烈岩爆虽少，但[[破坏最重]]', '中等以上岩爆占 [[38%]]'], notes='[gallery] chart_page')
    _research_layouts(d, F)
    _pro_layouts(d, F)
    d.section(3, variant='number', desc=['破裂演化机理与能量判据', '多源融合的前兆识别方法', '预警平台与工程示范'],
              notes='[gallery] section(i, variant="number") 大序号过渡页')
    _research_page(d)
    d.ending(notes='[gallery] ending() = 谢谢 (variant="band")')
    d.ending(variant='sky', notes='[gallery] ending(variant="sky")')
    d.ending(variant='split', notes='[gallery] ending(variant="split")')
    d.ending(variant='plaque', notes='[gallery] ending(variant="plaque")')
    d.ending(variant='midband', notes='[gallery] ending(variant="midband")')


def _pro_layouts(d, F):
    """Layouts from the 新参考 review (layouts_pro + parts): 风险 / 预期成果 / 课题分解 / 团队 / 负责人 / 背景四卡 /
    科研架构图 / 发展历程 / 获奖."""
    d.risks('风险分析与应对', statement='风险并非孤立存在，需建立[[风险预判]]与协同管控体系',
            risks=[{'name': '样本不足', 'summary': '极端岩爆样本稀缺', 'impact': '模型对强岩爆的识别能力不足，可能导致漏报', 'level': '高', 'topic': '课题二'},
                   {'name': '地质突变', 'summary': '断层带围岩参数突变', 'impact': '判据阈值失配，预警提前量下降', 'level': '中', 'topic': '课题一'},
                   {'name': '数据中断', 'summary': '光纤与微震链路受施工干扰', 'impact': '实时监测出现空窗，影响连续预警', 'level': '中', 'topic': '课题三'},
                   {'name': '推广受限', 'summary': '不同工程地质条件差异大', 'impact': '模型迁移需重新标定，推广周期延长', 'level': '低', 'topic': '课题三'}],
            conclusion='防控需坚持“机理约束 + 数据驱动”，兼顾[[可解释性]]与时效性',
            notes='[gallery] risks (cards: 级别 / 涉及课题 chips)')
    d.risks('风险分析（申报书版）', variant='tabs',
            risks=[{'name': '技术融合风险', 'label': '智能计算', 'summary': '特性预测精度不足', 'points': ['多源数据时空尺度不一致，融合误差累积'], 'level': '低', 'topic': '课题一'},
                   {'name': '核心工艺风险', 'label': '传感布设', 'summary': '深埋段安装难度大', 'points': ['高地应力区钻孔易塌孔，传感器存活率低'], 'level': '高', 'topic': '课题二'},
                   {'name': '数据安全风险', 'label': '平台运维', 'summary': '标准化程度不足', 'points': ['多单位数据共享存在权限与接口风险'], 'level': '中', 'topic': '课题三'},
                   {'name': '工程窗口风险', 'label': '现场试验', 'summary': '示范段工期受限', 'points': ['施工进度变化压缩现场试验窗口'], 'level': '中', 'topic': '课题三'},
                   {'name': '算力成本风险', 'label': '模型训练', 'summary': '三维反演计算量大', 'points': ['高精度数值反演耗时长，难以实时'], 'level': '中', 'topic': '课题一'},
                   {'name': '人才缺口风险', 'label': '团队稳定', 'summary': '交叉人才不足', 'points': ['岩石力学与人工智能复合型人才短缺'], 'level': '高', 'topic': '课题二'}],
            notes='[gallery] risks(variant="tabs") — level-colored chips')
    d.outcomes('预期成果', items=[{'title': '揭示灾变演化机理', 'text': '阐明深埋围岩能量积聚—释放的多尺度演化规律'},
                                 {'title': '建立预警判据体系', 'points': ['储耗比 + 位移速率双判据', '三级分级预警阈值']},
                                 {'title': '研发智能预警平台', 'text': '多源监测数据实时融合，预警提前量 ≥ 6 d'},
                                 {'title': '形成标准与示范', 'text': '编制技术指南 1 部，示范工程 2 处'}],
               center={'header': '预期成果', 'image': F['risk3d'], 'caption': '“机理 + 数据”双驱动的围岩灾变预警体系'},
               figures=[{'image': F['stress_field'], 'caption': '应力场反演'}, {'image': F['microseismic'], 'caption': '微震事件定位'},
                        {'image': F['curves'], 'caption': '位移—时间曲线'}],
               conclusion='成果将直接服务深埋长隧洞[[建设安全]]，支撑行业技术标准修订',
               notes='[gallery] outcomes (centre card + numbered outline pills + figure row)')
    d.topics('课题分解与关联', banner='系统集成与工程验证  ➜  [机理试验] → [智能预警] → [示范应用]',
             topics=[{'name': '课题1', 'title': '深埋围岩灾变机理与[[能量判据]]', 'image': F['stress_field']},
                     {'name': '课题2', 'title': '多源感知与前兆特征智能识别', 'image': F['microseismic']},
                     {'name': '课题3', 'title': '分级预警模型与云平台', 'image': F['network']},
                     {'name': '课题4', 'title': '示范工程与[[现场验证]]', 'image': F['route_map']}],
             links={(1, 2): ('机理约束', '数据反馈'), (1, 3): '判据输入', (2, 3): '特征支撑', (3, 4): '迭代优化闭环'},
             side_left='提供机理输入', side_right='依赖现场反馈',
             conclusion='覆盖从基础机理到工程应用的全链条，四个课题相互支撑、迭代闭环',
             notes='[gallery] topics (2×2 课题 panels + labeled links + side brackets)')
    d.lab_team('研究团队与人员构成', headline='固定人员共计[[46]]人，其中40岁以下[[29]]人',
           stats=[('28', '人', '研究人员'), ('12', '人', '技术人员'), ('6', '人', '管理人员')],
           leader={'name': '张三', 'label': '项目负责人', 'titles': ['教授', '博士生导师'],
                   'points': ['主持国家自然科学基金重点项目、国家重点研发计划课题等 [[8]] 项', '获省部级科技进步一等奖 [[2]] 项']},
           members=[{'name': '李四', 'role': '课题1负责人'}, {'name': '王五', 'role': '课题2负责人'},
                    {'name': '赵六', 'role': '课题3负责人'}, {'name': '钱七', 'role': '课题4负责人'}],
           columns=['职务', '姓名', '职称', '工作单位'],
           rows=[['负责人', '张三', '教授', '示例大学'], ['副负责人', '李四', '教授', '示例大学'], ['骨干', '王五', '研究员', '示例工程研究院'],
                 ['骨干', '赵六', '副教授', '示例大学'], ['骨干', '钱七', '高级工程师', '示例水利水电工程局'], ['骨干', '孙八', '副研究员', '示例工程研究院'],
                 ['骨干', '周九', '讲师', '示例理工大学'], ['骨干', '吴十', '工程师', '示例水利水电工程局']],
           highlight_rows=(0, 1), notes='[gallery] lab_team (stats + leader + photo row | roster table)')
    d.leader('项目负责人', name='张三', honors=['项目负责人', '省重点实验室主任', '国家优青', '省杰青'],
             points=['主持国家重点研发计划课题、国家自然科学基金重点项目等 [[12]] 项',
                     '获省部级科技进步[[一等奖 2 项]]、二等奖 3 项', '发表 SCI 论文 [[60]] 余篇，授权发明专利 [[25]] 项'],
             awards=[{'title': '省科技进步奖', 'level': '一等奖'}, {'title': '中国岩石力学学会', 'level': '科技进步奖'},
                     {'title': '水力发电科学技术', 'level': '二等奖'}],
             groups=[{'name': '课题一团队', 'members': [{'name': '李四', 'role': '（教授）'}, {'name': '王五', 'role': '（研究员）'}]},
                     {'name': '课题二团队', 'members': [{'name': '赵六', 'role': '（副教授）'}, {'name': '钱七', 'role': '（高工）'}]},
                     {'name': '课题三团队', 'members': [{'name': '孙八', 'role': '（副研究员）'}, {'name': '周九', 'role': '（讲师）'}]}],
             notes='[gallery] leader (photo + honours + laurels + sub-team groups)')
    d.need_cards('研究背景', statement='深埋隧洞建设向[[超深、超长]]发展，围岩灾变已成为制约工程安全的核心瓶颈',
                 cards=[{'tab': '科学挑战', 'image': F['stress_field'], 'caption': '机理认识不足', 'head': '能量积聚—释放', 'sub': '多尺度演化规律不清', 'foot': '高地应力下围岩破裂演化难以刻画'},
                        {'tab': '国家需求', 'image': F['route_map'], 'caption': '重大工程密集', 'head': '引调水与交通隧洞', 'sub': '埋深超过 2000 m', 'foot': '工程安全直接关系民生与投资效益'},
                        {'tab': '国际竞争', 'image': F['network'], 'caption': '智能预警前沿', 'head': '感知与智能融合', 'sub': '抢占技术制高点', 'foot': '高端监测装备与算法亟需自主可控'},
                        {'tab': '技术发展', 'image': F['microseismic'], 'caption': '多源实时感知', 'head': '光纤 + 微震', 'sub': '海量监测数据', 'foot': '为数据驱动的智能预警提供基础'}],
                 converge={'left': '揭示灾变机理，为预警提供物理约束', 'right': '融合多源数据，提升预警时效与精度', 'center': '科学机遇',
                           'conclusion': '从根本上提升深埋隧洞围岩灾变的预测预警能力'},
                 notes='[gallery] need_cards (tab cards + chevrons + V-convergence)')
    d.framework('研究框架', banner='构建“机理—感知—预警—应用”一体化研究框架', center='围岩灾变\n智能预警',
                ring=[{'title': '灾变机理', 'subs': ['能量', '破裂']}, {'title': '多源感知', 'subs': ['光纤', '微震']},
                      {'title': '预警模型', 'subs': ['判据', '阈值']}, {'title': '工程应用', 'subs': ['示范', '标准']}],
                accents=['理', '感', '判', '用'], arcs=['机理约束', '数据驱动', '分级预警', '反馈修正'],
                left={'circles': ['基础\n理论', '试验\n方法'], 'boxes': [{'title': '理论层', 'text': '能量判据与演化规律'}, {'title': '方法层', 'text': '真三轴 + 数值反演'}]},
                right={'circles': ['平台\n研发', '工程\n示范'], 'boxes': [{'title': '技术层', 'text': '多源融合与智能识别'}, {'title': '应用层', 'text': '分级预警与决策支持'}]},
                bottom=['现场监测', '数据融合', '智能预警', '风险决策', '工程反馈'], palette='teal',
                notes='[gallery] framework (科研架构图 hub, palette="teal")')
    d.history('团队发展历程', banner='团队发展历程', subtitle='从“经验判断”到“机理 + 数据”智能预警的演进',
              events=[{'year': '2008年', 'title': '岩石力学实验平台建成', 'stars': 1, 'image': F['curves'], 'key': True},
                      {'year': '2012年', 'title': '真三轴试验系统研制', 'stars': 2, 'image': F['stress_field']},
                      {'year': '2016年', 'title': '微震监测网络部署', 'stars': 2, 'image': F['microseismic']},
                      {'year': '2020年', 'title': '智能预警模型突破', 'stars': 3, 'image': F['network'], 'key': True},
                      {'year': '2025年', 'title': '预警云平台示范应用', 'stars': 2, 'image': F['route_map']}],
              conclusion='逐步形成“试验—监测—预警—应用”完整技术体系',
              notes='[gallery] history (发展历程: year pills + ★ + photos)')
    d.awards('科研团队获奖', banner='获得省部级以上科技奖励共[[12]]项，其中第一完成单位[[8]]项',
             columns=['序号', '奖励名称', '项目名称', '获奖年份'],
             rows=[['1', '省科技进步奖 一等奖', '深埋隧洞围岩稳定控制关键技术', '2023'],
                   ['2', '中国岩石力学与工程学会科技进步奖 一等奖', '岩爆智能预警理论与方法', '2022'],
                   ['3', '水力发电科学技术奖 二等奖', '引调水工程隧洞安全监测技术', '2021']],
             highlight_rows=(0,), certificates=['荣誉证书', 'CERTIFICATE', '荣誉证书', 'CERTIFICATE', '荣誉证书'],
             notes='[gallery] awards (banner + table + certificate shelf)')
    d.risks('风险清单与应对措施', variant='rows',
            risks=[{'name': '样本不足', 'topic': '课题二', 'impact': '强岩爆样本稀缺，模型对极端事件识别能力不足', 'level': '高', 'measure': '联合多工程共建样本库，引入物理约束数据增强'},
                   {'name': '地质突变', 'topic': '课题一', 'impact': '断层带参数突变导致判据阈值失配', 'level': '中', 'measure': '超前地质预报联动，阈值动态修正'},
                   {'name': '数据中断', 'topic': '课题三', 'impact': '施工干扰导致监测链路中断', 'level': '中', 'measure': '冗余链路 + 边缘缓存，断点续传'},
                   {'name': '推广受限', 'topic': '课题三', 'impact': '不同地质条件下模型迁移困难', 'level': '低', 'measure': '迁移学习 + 少样本标定流程'}],
            notes='[gallery] risks(variant="rows") — level badges + 应对措施')


# ------------------------------------------------------------------ (2) style showcase
def _showcase(d, skin):
    """The same 9-page sample talk in one skin (one 3x3 sheet per skin)."""
    from .theme import SKINS
    F = _figs()
    _meta(d, project=f'国家自然科学基金 · 结题汇报（{SKINS[skin]["label"]} · 示例）')
    lab = f'[gallery] {skin} · '
    d.cover(notes=lab + 'cover')
    d.agenda(descs=['工程需求与科学问题', '技术路线与关键方法', '主要成果与创新点', '结论与应用前景'], notes=lab + 'agenda')
    d.cur_section = 0
    _p_need(d, F, notes=lab + 'panels')
    _p_problems(d, F, notes=lab + 'problems')
    d.cur_section = 1
    d._sub_counter[1] = 0
    _p_route(d, F, notes=lab + 'route')
    _p_method(d, F, notes=lab + 'figure_text')
    d.cur_section = 2
    d._sub_counter[2] = 0
    _p_results(d, F, notes=lab + 'panels + table')
    _p_innov(d, F, notes=lab + 'innovations')
    d.ending(notes=lab + 'ending (谢谢)')


def build_showcase(out_dir, sheets_dir=None, skins=None):
    """Build one rendered sample deck per skin; with sheets_dir write labeled 3x3 sheets 'style_<skin>.jpg'."""
    from .theme import SKINS, SKIN_THEMES
    out = []
    for skin in (skins or SKINS):
        d = Deck(SKIN_THEMES[skin], skin=skin)
        _showcase(d, skin)
        p = d.save(os.path.join(out_dir, f'showcase_{skin}.pptx'), render=True, check=False)
        if sheets_dir:
            for s in export_sheets(p, os.path.join(sheets_dir, f'style_{skin}'), per=9):
                final = s.replace('_1.jpg', '.jpg')
                if os.path.exists(final):
                    os.remove(final)
                os.replace(s, final)
                out.append(final)
    return out


def _research_layouts(d, F):
    """Layouts from the 2024–2026 tutorial-image review (layouts_rs + diagram helpers)."""
    d.band_cards('研究目标与研究内容', headline='揭示深埋围岩灾变机理，构建“机理+数据”双驱动的[[智能预警]]体系',
                 text='围绕“看得见、判得准、报得及”三个环节，设置三项相互衔接的研究内容。',
                 kpis=[{'value': '6', 'unit': 'd', 'label': '预警提前量目标'}, {'value': '90', 'unit': '%',
                                                                          'label': '前兆识别准确率'}],
                 cards=[{'title': '研究内容一', 'icon': 'activity', 'tag': '机理',
                         'points': ['真三轴加卸载试验', '多尺度破裂演化规律', '岩爆[[能量释放]]判据']},
                        {'title': '研究内容二', 'icon': 'radar', 'tag': '感知',
                         'points': ['光纤 + 微震联合感知', '多源数据实时融合', '前兆特征[[智能识别]]']},
                        {'title': '研究内容三', 'icon': 'bell-ring', 'tag': '预警',
                         'points': ['风险分级与阈值', '预警云平台研发', '示范段[[工程验证]]']}],
                 notes='[gallery] band_cards (拦腰: band headline + KPIs, overlapping cards)')
    d.band_cards('能量判据与预警模型', headline='以[[弹性应变能]]储耗比判别岩爆倾向，以位移速率触发分级预警',
                 pills=['储耗比', '应力集中', '位移速率', '三级阈值'],
                 cards=[{'title': '岩爆倾向判据', 'formula': r'W_{et}=\dfrac{\phi_{sp}}{\phi_{st}}\ \geq\ 2.0',
                         'symbols': [(r'W_{et}', '储耗比'), (r'\phi_{sp}', '弹性应变能'), (r'\phi_{st}', '耗散能')]},
                        {'title': '分级预警准则', 'formula': r'v=\dfrac{\Delta u}{\Delta t},\quad v\geq v_{c}',
                         'symbols': [(r'u', '累计位移'), (r'v', '位移速率'), (r'v_c', '预警阈值')]}],
                 card_style='topline', notes='[gallery] band_cards(card_style="topline") with formula cards')
    d.architecture('智能预警系统总体架构', top='目标：深埋长隧洞围岩灾变“可感知、可诊断、可预警”',
                   tiers=[{'label': '应用层', 'items': [{'text': '分级预警', 'icon': 'bell-ring'},
                                                        {'text': '处置预案', 'icon': 'clipboard-list'},
                                                        {'text': '辅助决策', 'icon': 'sliders-horizontal'},
                                                        {'text': '移动巡检', 'icon': 'smartphone'}]},
                          {'label': '诊断层', 'groups': [{'title': '机理模型', 'items': ['能量判据', '数值反演']},
                                                        {'title': '数据模型', 'items': ['前兆识别', '趋势预测', '异常检测']},
                                                        {'title': '融合决策', 'items': ['风险评估', '阈值更新']}]},
                          {'label': '传输层', 'items': ['边缘网关', '5G 回传', '时间同步', '云边协同']},
                          {'label': '感知层', 'items': ['微震', '光纤应变', '多点位移', '应力', '渗压', '视频']}],
                   sides=[{'label': '标准规范体系', 'side': 'left'}, {'label': '安全保障体系', 'side': 'right'}],
                   notes='[gallery] architecture (tiers + groups + side columns)')
    d.versus('与传统预警方法对比', statement='本文方法在[[预警提前量]]与[[识别准确率]]上全面优于传统方法（示例数据）',
             ours={'name': '本文双驱动预警', 'sub': '机理判据 + 深度学习', 'badge': '综合最优',
                   'rows': [('识别准确率', '92.4%'), ('预警提前量', '6.0 d'), ('误报率', '4.1%'), ('响应时间', '< 1 min')]},
             others=[{'name': '经验判识', 'sub': '人工巡检 + 规范',
                      'rows': [('识别准确率', '68.0%'), ('预警提前量', '1.5 d'), ('误报率', '21.3%'), ('响应时间', '数小时')]},
                     {'name': '单一数据模型', 'sub': '机器学习',
                      'rows': [('识别准确率', '84.7%'), ('预警提前量', '3.2 d'), ('误报率', '9.6%'), ('响应时间', '分钟级')]}],
             notes='[gallery] versus (ours highlighted vs gray + VS)')
    d.process_matrix('研究技术路线（矩阵式）', steps=['试验研究', '机理揭示', '模型构建', '平台研发', '工程示范'],
                     rows=['理论', '方法', '应用'], phases=[('基础研究', 0, 1), ('关键技术', 2, 3), ('示范应用', 4, 4)],
                     cells=[[{'title': '真三轴试验', 'text': '120 组加卸载'}, {'title': '破裂演化', 'text': '能量释放规律'},
                             None, None, None],
                            [None, {'title': '能量判据', 'text': '储耗比阈值'}, {'title': '识别模型', 'text': '时频 + 深度学习'},
                             {'title': '预警算法', 'text': '三级阈值联动'}, None],
                            [None, None, None, {'title': '预警云平台', 'text': '云边协同'},
                             {'title': '示范工程', 'text': '32.6 km 隧洞'}]],
                     arrows=[(0, 0, 0, 1), (0, 1, 1, 1), (1, 1, 1, 2), (1, 2, 1, 3), (1, 3, 2, 3), (2, 3, 2, 4)],
                     notes_row=['数据可靠', '机理清楚', '模型可用', '平台可靠', '工程见效'],
                     notes='[gallery] process_matrix (phases + chevrons + cells + arrows)')
    d.cycle('“感知—诊断—预警—处置”闭环', center='闭环\n预警', center_sub='分钟级响应',
            steps=[{'label': '实时感知', 'title': '实时感知', 'points': ['微震 + 光纤[[全断面]]覆盖', '1 Hz 连续采集']},
                   {'label': '智能诊断', 'title': '智能诊断', 'points': ['前兆特征识别', '风险等级评估']},
                   {'label': '分级预警', 'title': '分级预警', 'points': ['蓝 / 黄 / 红三级阈值', '多渠道[[即时推送]]']},
                   {'label': '联动处置', 'title': '联动处置', 'points': ['处置预案自动匹配', '效果反馈与模型更新']}],
            statement='以[[智能诊断]]为核心，打通感知、诊断、预警与处置环节', notes='[gallery] cycle (closed loop)')
    nav0, d.nav = d.nav, 'bottom'
    d.bands('国内外研究现状', bands=[
        {'word': '国外', 'label': '国外研究现状', 'cols': [
            {'title': '灾变机理', 'points': ['真三轴试验揭示破裂模式', '能量判据用于岩爆分级']},
            {'title': '监测预警', 'points': ['微震监测系统成熟', '数据驱动预警逐步兴起']}]},
        {'word': '国内', 'label': '国内研究现状', 'cols': [
            {'title': '灾变机理', 'points': ['深埋工程岩爆案例丰富', '多尺度机理认识不足']},
            {'title': '监测预警', 'points': ['多源监测广泛应用', '预警阈值依赖[[经验]]']}]}],
        conclusion='不足：机理与数据割裂，缺少面向深埋长隧洞的实时智能预警方法',
        notes='[gallery] bands (国内外研究现状) + nav="bottom"')
    d.nav = nav0
    d.tree('课题分解与任务设置', root='深埋长隧洞围岩灾变机理与智能预警', root_sub='项目总体目标',
           children=[{'name': '课题一', 'sub': '灾变机理', 'children': ['真三轴试验', '破裂演化模型', '能量判据']},
                     {'name': '课题二', 'sub': '智能感知', 'children': ['光纤微震联合', '数据融合']},
                     {'name': '课题三', 'sub': '前兆识别', 'children': ['特征提取', '识别模型', '模型验证']},
                     {'name': '课题四', 'sub': '预警平台', 'children': [{'name': '平台研发', 'sub': '云边协同'},
                                                                       {'name': '工程示范', 'sub': '32.6 km'}]}],
           notes='[gallery] tree (课题分解)')
    d.panels('研究进展与指标完成情况', statement='主要考核指标[[全部达成]]，三项指标超额完成（示例）', panels=[
        {'header': '核心指标', 'stat_grid': [{'label': '监测里程', 'value': 32.6, 'unit': 'km'},
                                            {'label': '采样频率', 'value': 1, 'unit': 'Hz'},
                                            {'label': '预警提前量', 'value': 6, 'unit': 'd', 'color': 'emph'},
                                            {'label': '识别准确率', 'value': 92.4, 'unit': '%'},
                                            {'label': '误报率', 'value': 4.1, 'unit': '%'},
                                            {'label': '成功预警', 'value': 12, 'unit': '次'}],
         'points': ['示范段连续运行 [[18 个月]]', '平台在 2 个工程推广应用']},
        {'header': '任务完成度', 'bars': [{'label': '机理试验', 'value': 100}, {'label': '感知网络', 'value': 100},
                                          {'label': '识别模型', 'value': 95}, {'label': '预警平台', 'value': 90},
                                          {'label': '工程示范', 'value': 80}], 'bars_highlight': 4,
         'points': ['工程示范持续推进，年底完成验收']}],
        notes='[gallery] panels with stat_grid / bars')


def _research_page(d):
    """formula() + zoom_inset() + callout() + named method colors + nav='mini' on one page."""
    from . import academic as A
    from . import samples
    from .canvas import Box
    d.add_colors(ma='0072BD', mb='C00000')
    nav0, d.nav = d.nav, 'mini'
    cv = d.page('力学模型与监测验证（科研图文）')
    d.nav = nav0
    c = cv.content
    y = c.y + A.statement(cv, c.x, c.y, c.w, '公式按真实字号排版；原图[[局部放大]]代替红框；方法名与原图同色'
                                            '（{{ma|监测点A}} / {{mb|本文模型}}）') + 0.17
    L = Box(c.x, y, 5.6, c.b - y)
    ip = A.panel(cv, L, '围岩应力与能量判据')
    A.formula(cv, Box(ip.x, ip.y + 0.1, ip.w, 0.95),
              r'\sigma_\theta=\dfrac{\sigma_v+\sigma_h}{2}\left(1+\dfrac{a^2}{r^2}\right)-'
              r'\dfrac{\sigma_v-\sigma_h}{2}\left(1+\dfrac{3a^4}{r^4}\right)\cos2\theta', num=1, size=17,
              note='a 洞室半径 · r、θ 极坐标 · σv、σh 竖向与水平地应力')
    A.formula(cv, Box(ip.x, ip.y + 1.45, ip.w, 0.85), r'W_{et}=\dfrac{\phi_{sp}}{\phi_{st}}\ \geq\ 2.0', num=2, size=19,
              note='弹性能储耗比，≥2.0 判为岩爆倾向')
    A.arrow_list(cv, Box(ip.x, ip.y + 2.55, ip.w, ip.h - 2.55),
                 ['开挖后拱顶切向应力集中系数可达 [[2.4]]', '储耗比超过阈值区段列为[[重点监测]]',
                  '判据与监测数据联合校核'], size=13)
    fig = samples.curves()
    fb = Box(L.r + 0.3, y, 4.1, 3.05)
    pic = A.figure(cv, fb, fig, '图7 位移监测（示例）')
    A.zoom_inset(cv, fig, pic, (0.62, 0.38, 0.95, 0.9), Box(fb.r + 0.4, y + 0.38, c.r - fb.r - 0.4, 2.35),
                 label='预警窗口')
    A.callout(cv, fb.x + fb.w * 0.36, fb.y + fb.h * 0.5, fb.x + 0.1, fb.b + 0.3, 3.6, 0.5, '曲线拐点：前兆特征出现')
    cv.notes('[gallery] formula / zoom_inset / callout / nav="mini" / add_colors')


def _modern(d):
    """Modern style (日常汇报 / 课件 / 路演) catalog with a realistic sample topic: 智慧水务平台年度工作汇报."""
    F = _figs()
    d.meta(title='智慧水务监测平台\n2026 年度工作汇报', subtitle='数据驱动 · 安全运行 · 降本增效', en='ANNUAL WORK REPORT',
           author='张三', org='示例水务集团 信息中心', date='2026.12')
    d.set_sections(['年度概况', '重点工作', '成效数据', '明年计划'])
    d.cover(notes='[gallery] cover(variant="left")')
    d.cover(variant='center', notes='[gallery] cover(variant="center")')
    d.cover(variant='split', notes='[gallery] cover(variant="split")')
    d.agenda(notes='[gallery] agenda(variant="list")')
    d.agenda(variant='cards', notes='[gallery] agenda(variant="cards")')
    d.section(1, desc='全年平台建设与运行情况', notes='[gallery] section(i, variant="bignum")')
    d.section(2, variant='nav', notes='[gallery] section(i, variant="nav")')
    d.cards('三大建设目标', [{'icon': 'radar', 'title': '全域感知', 'text': '管网压力、流量、水质在线监测覆盖率达 95%'},
                         {'icon': 'cpu', 'title': '智能诊断', 'text': '漏损、爆管、水质异常自动识别与定位'},
                         {'icon': 'shield-check', 'title': '安全运行', 'text': '重大事件响应时间缩短至 15 分钟以内'}],
            summary='以“感知—诊断—处置”闭环提升供水安全保障能力', notes='[gallery] cards(variant="icon_top", summary=...)')
    d.cards('六项重点任务', [{'title': t_, 'text': x} for t_, x in (
        ('监测点扩容', '新增压力监测点 320 个'), ('平台升级', '完成 3.0 版本上线'), ('漏损治理', '分区计量覆盖 48 个片区'),
        ('水质预警', '建成 12 项指标预警模型'), ('应急演练', '组织联合演练 6 次'), ('人才培养', '完成专项培训 860 人次'))],
        variant='number', notes='[gallery] cards(variant="number")')
    d.cards('平台核心能力', [{'icon': ic, 'title': t_, 'text': x} for ic, t_, x in (
        ('activity', '实时监测', '秒级采集，分钟级汇聚'), ('scan-search', '异常识别', '多源数据融合诊断'),
        ('map', '一张图调度', 'GIS 管网与设备联动'), ('bell-ring', '分级预警', '短信、App 多渠道推送'))],
        variant='icon_left', notes='[gallery] cards(variant="icon_left")')
    d.bullets('年度工作亮点', ['**漏损率下降**：全市综合漏损率由 12.6% 降至 9.8%',
                             '**平台升级**：3.0 版本上线，接入设备 1.2 万台',
                             '**应急提速**：爆管平均处置时长缩短 38%'],
              lead='2026 年围绕“安全、高效、智慧”三条主线推进：', notes='[gallery] bullets(style="rows")')
    d.text_image('分区计量与漏损定位', heading='把城市管网划分为 48 个独立计量片区',
                 text='通过片区进出水量平衡与夜间最小流量分析，快速锁定漏损片区，再以压力波动特征缩小排查范围。',
                 image=F['network'], points=['平均定位时间由 3 天缩短至 6 小时', '全年发现隐蔽漏点 214 处'],
                 notes='[gallery] text_image')
    d.kpis('年度核心指标', [{'value': '9.8', 'unit': '%', 'label': '综合漏损率', 'desc': '同比下降 2.8 个百分点'},
                          {'value': '1.2', 'unit': '万台', 'label': '接入设备'},
                          {'value': '15', 'unit': 'min', 'label': '应急响应'}], notes='[gallery] kpis(variant="cards")')
    d.kpis('监测覆盖率', [{'value': '95%', 'label': '压力监测'}, {'value': '88%', 'label': '流量计量'},
                          {'value': '76%', 'label': '水质在线'}], variant='rings', notes='[gallery] kpis(variant="rings")')
    d.chart_page('各片区漏损率', {'kind': 'bar', 'categories': ['东区', '西区', '南区', '北区', '开发区'],
                                'series': {'漏损率': [8.2, 10.5, 9.1, 12.3, 7.4]}, 'highlight': 3, 'number_format': '0.0'},
                 points=['北区管龄长，漏损率[[最高]]', '开发区新建管网表现最好'], chart_title='2026 年各片区漏损率',
                 unit='单位：%', notes='[gallery] chart_page(kind="bar")')
    d.table_page('重点项目进度', ['项目', '计划完成', '实际完成', '完成率'],
                 [['监测点扩容', '300 个', '320 个', '107%'], ['分区计量', '45 个片区', '48 个片区', '107%'],
                  ['平台 3.0 升级', '12 月', '11 月', '提前'], ['水质预警模型', '10 项', '12 项', '120%']],
                 highlight_col=3, notes='[gallery] table_page(highlight_col=3)')
    ev = [{'date': y, 'title': t_, 'text': x} for y, t_, x in (
        ('2023', '平台立项', '完成顶层设计'), ('2024', '一期上线', '接入主干管网'), ('2025', '二期扩容', '覆盖全部片区'),
        ('2026', '3.0 升级', '智能诊断上线'))]
    d.timeline('平台发展历程', ev, highlight=3, notes='[gallery] timeline(variant="axis")')
    d.timeline('三年建设阶段', ev[1:], variant='cards', highlight=2, notes='[gallery] timeline(variant="cards")')
    d.timeline('能力演进', [{'date': v, 'title': t_, 'text': x} for v, t_, x in (
        ('1.0', '看得见', '数据采集与展示'), ('2.0', '管得住', '告警与工单联动'), ('3.0', '算得准', '智能诊断与预测'))],
        variant='rise', notes='[gallery] timeline(variant="rise")')
    st = [{'title': '异常发现', 'text': '监测数据自动告警', 'icon': 'bell-ring'},
          {'title': '智能诊断', 'text': '定位漏损片区', 'icon': 'cpu'},
          {'title': '现场处置', 'text': '工单派发与闭环', 'icon': 'wrench'}]
    d.process('爆管处置流程', st, notes='[gallery] process(variant="chevron")')
    d.process('漏损治理四步法', [{'title': t_, 'text': x} for t_, x in (
        ('分区', '划分计量片区'), ('计量', '进出水量平衡'), ('定位', '压力特征排查'), ('修复', '修复与复核'))],
        variant='circles', notes='[gallery] process(variant="circles")')
    d.process('数据治理流程', st, variant='cards', notes='[gallery] process(variant="cards")')
    d.hub('智慧水务一张图', '一张图\n调度', [{'icon': 'radar', 'title': '监测', 'text': '压力、流量、水质'},
                                       {'icon': 'cpu', 'title': '诊断', 'text': '漏损与爆管识别'},
                                       {'icon': 'map', 'title': '调度', 'text': '泵站与阀门联动'},
                                       {'icon': 'shield-check', 'title': '应急', 'text': '预案与演练'}],
          notes='[gallery] hub')
    d.compare('人工巡检 vs 智能监测', {'title': '人工巡检', 'points': ['周期长，隐蔽漏点难发现', '依赖经验，定位慢']},
              {'title': '智能监测', 'points': ['7×24 小时在线感知', '自动定位，处置时长缩短 38%']},
              notes='[gallery] compare(variant="cards")')
    d.matrix('问题优先级', [{'title': '北区老旧管网', 'text': '高风险 · 高影响', 'icon': 'triangle-alert'},
                         {'title': '水质在线不足', 'text': '中风险 · 高影响', 'icon': 'flask-conical'},
                         {'title': '设备掉线', 'text': '高频 · 低影响', 'icon': 'wifi-off'},
                         {'title': '数据标准不一', 'text': '低风险 · 中影响', 'icon': 'database'}],
             x_label='发生频率 →', y_label='影响程度', highlight=0, notes='[gallery] matrix')
    d.pyramid('数据价值金字塔', [{'title': '决策', 'text': '调度与投资决策'}, {'title': '诊断', 'text': '漏损与风险识别'},
                               {'title': '数据', 'text': '1.2 万台设备实时采集'}], notes='[gallery] pyramid(variant="pyramid")')
    d.pyramid('平台技术架构', [{'title': '应用层', 'items': ['一张图调度', '移动巡检']},
                             {'title': '平台层', 'items': ['数据中台', '模型服务']},
                             {'title': '感知层', 'items': ['压力流量', '水质在线']}], variant='layers',
              notes='[gallery] pyramid(variant="layers")')
    d.band('三项保障措施', '以安全运行为底线，以数据驱动为核心', [
        {'icon': 'users', 'title': '组织保障', 'text': '成立专项工作组'}, {'icon': 'coins', 'title': '资金保障', 'text': '年度投入 2800 万元'},
        {'icon': 'graduation-cap', 'title': '人才保障', 'text': '专项培训 860 人次'}], notes='[gallery] band')
    d.split('明年工作计划', [{'title': t_, 'points': p_, 'tag': '0' + str(k + 1)} for k, (t_, p_) in enumerate((
        ('一季度', ['完成北区管网改造', '新增监测点 150 个']), ('二至三季度', ['水质预警模型扩展', '移动巡检全面推广']),
        ('四季度', ['平台 3.5 版本上线', '年度考核与总结'])))], arrows=True, notes='[gallery] split(arrows=True)')
    d.team('项目团队', [{'name': n, 'role': r, 'desc': x} for n, r, x in (
        ('张三', '项目负责人', '统筹平台建设'), ('李四', '技术负责人', '数据与模型'), ('王五', '运维负责人', '设备接入与运维'))],
        notes='[gallery] team')
    d.quote('让每一滴水都可感知、可计量、可追溯。', source='平台建设理念', notes='[gallery] quote')
    d.ending(notes='[gallery] ending() = 谢谢 (variant="thanks")')


def build_gallery(style='defense', out=None, theme=None, sheets_prefix=None):
    """Build the gallery deck, render it, and (optionally) write labeled 3x3 sheets '<prefix>_N.jpg'."""
    theme = theme or ('nsfc_blue' if style == 'defense' else 'academic_blue')
    d = Deck(theme)
    (_defense if style == 'defense' else _modern)(d)
    out = out or os.path.abspath(f'gallery_{style}.pptx')
    path = d.save(out, render=True)
    if sheets_prefix:
        export_sheets(path, sheets_prefix)
    return path


def export_sheets(pptx, prefix, per=9, cols=3, thumb_w=620):
    """Labeled contact sheets from a rendered gallery deck (labels = '[gallery] ...' speaker notes)."""
    import glob
    from pptx import Presentation
    from . import render as R
    prs = Presentation(pptx)
    labels = []
    for i, sl in enumerate(prs.slides, 1):
        n = ''
        if sl.has_notes_slide:
            n = sl.notes_slide.notes_text_frame.text.strip()
        labels.append(f'{i:02d} ' + (n.replace('[gallery]', '').strip() if n.startswith('[gallery]') else ''))
    prev = os.path.splitext(os.path.abspath(pptx))[0] + '_preview'
    pngs = sorted(glob.glob(os.path.join(prev, 'slide*.png')))
    outs = []
    for k in range(0, len(pngs), per):
        o = f'{prefix}_{k // per + 1}.jpg'
        R.contact_sheet(pngs[k:k + per], o, cols=cols, thumb_w=thumb_w, labels=labels[k:k + per])
        outs.append(o)
    return outs
