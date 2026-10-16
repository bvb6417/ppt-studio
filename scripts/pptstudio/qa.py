"""Static QA on a saved .pptx (no PowerPoint needed):
  * off-slide content, overlapping text boxes, tiny fonts, estimated text overflow, font count.
Shapes named 'bg:*' / 'deco:*' are decorative and exempt from overlap / off-slide rules.
For exact overflow use render.check_overflow() (PowerPoint layout engine).
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Emu

from . import fonts as F

EMU = 914400


def _boxes(slide):
    out = []
    for sh in slide.shapes:
        try:
            x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
        except Exception:
            continue
        out.append((sh, x, y, w, h))
    return out


def _runs(sh):
    if not getattr(sh, 'has_text_frame', False) or not sh.has_text_frame:
        return []
    rs = []
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            rs.append(r)
    return rs


def check(path, theme=None, min_font=10.0):
    prs = Presentation(path)
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    issues = []
    fonts_used = set()
    for si, slide in enumerate(prs.slides, 1):
        boxes = _boxes(slide)
        texts = []
        for sh, x, y, w, h in boxes:
            nm = sh.name or ''
            deco = nm.startswith(('bg:', 'deco:'))
            if not deco and (x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02):
                issues.append((si, 'off-slide', f'{nm} at ({x:.2f},{y:.2f},{w:.2f},{h:.2f})'))
            rs = _runs(sh)
            if rs and not deco:
                texts.append((sh, x, y, w, h))
            for r in rs:
                rPr = r._r.rPr
                if rPr is not None:
                    for tag in ('latin', 'ea'):
                        e = rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag)
                        if e is not None:
                            fonts_used.add(e.get('typeface'))
                    sz = rPr.get('sz')
                    if sz and int(sz) / 100 < min_font and not deco and r.text.strip() and \
                            not nm.startswith(('txt:source', 'txt:title_en', 'txt:kpilabel')):
                        issues.append((si, 'tiny-font', f'{nm}: {int(sz) / 100:.1f}pt "{r.text[:20]}"'))
            # estimated overflow for wrapped text boxes
            if rs and not deco and sh.text_frame.word_wrap is not False:
                try:
                    est = _estimate_height(sh)
                    if est and est > h * 72 + 3:
                        issues.append((si, 'overflow?', f'{nm}: needs ~{est / 72:.2f}in, box {h:.2f}in'))
                except Exception:
                    pass
        # overlapping text boxes
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i], texts[j]
                ix = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                iy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
                if ix > 0.08 and iy > 0.08:
                    area = ix * iy
                    small = min(a[3] * a[4], b[3] * b[4])
                    if small > 0 and area / small > 0.25:
                        issues.append((si, 'text-overlap', f'{a[0].name} x {b[0].name}'))
    return dict(path=path, slides=len(prs.slides), issues=issues, fonts=sorted(f for f in fonts_used if f))


def _estimate_height(sh):
    tf = sh.text_frame
    bp = tf._txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr')
    lI = int(bp.get('lIns', 91440)) / 12700
    rI = int(bp.get('rIns', 91440)) / 12700
    tI = int(bp.get('tIns', 45720)) / 12700
    bI = int(bp.get('bIns', 45720)) / 12700
    width = sh.width / 12700 - lI - rI
    total = tI + bI
    for p in tf.paragraphs:
        size, ea, lat, bold = 18.0, 'Microsoft YaHei', None, False
        txt = ''.join(r.text for r in p.runs)
        if p.runs:
            rPr = p.runs[0]._r.rPr
            if rPr is not None:
                if rPr.get('sz'):
                    size = int(rPr.get('sz')) / 100
                bold = rPr.get('b') == '1'
                A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
                if rPr.find(A + 'ea') is not None:
                    ea = rPr.find(A + 'ea').get('typeface')
                if rPr.find(A + 'latin') is not None:
                    lat = rPr.find(A + 'latin').get('typeface')
        pPr = p._p.pPr
        ls = 1.0
        sa = 0
        if pPr is not None:
            A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
            e = pPr.find(A + 'lnSpc/' + A + 'spcPct')
            if e is not None:
                ls = int(e.get('val')) / 100000
            e = pPr.find(A + 'spcAft/' + A + 'spcPts')
            if e is not None:
                sa = int(e.get('val')) / 100
        h, n, _ = F.layout_height([txt], width, size, ea, lat, bold, ls)
        total += h + sa
    return total


def format_report(rep):
    lines = [f"QA {rep['path']}: {rep['slides']} slides, fonts: {', '.join(rep['fonts'])}"]
    for si, kind, msg in rep['issues']:
        lines.append(f'  slide {si:02d} [{kind}] {msg}')
    if not rep['issues']:
        lines.append('  no static issues found')
    return '\n'.join(lines)
